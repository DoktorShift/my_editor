# SPDX-FileCopyrightText: 2026 rinbal
# SPDX-License-Identifier: AGPL-3.0-or-later
"""An ordinary workbook from Excel (Windows, Mac, the web), Google Sheets,
LibreOffice or Numbers is edited and saved into its own file, with
everything those apps put into it kept or moved along correctly.

Each test edits a cell, inserts a row, renames a sheet and saves, then
reads the package as Excel would: the zip, content types,
relationships, the order of every sheet's and the workbook's elements,
and what was kept (byte for byte where nothing in it named a cell)."""

import re
import zipfile

import openpyxl
import pytest
from lxml import etree

import sheet_formulas as sf
import sheet_io
from sheet_io import load_workbook
from tests import xlsx_producers as producers

# openpyxl warns about the parts it does not read; MyEditor keeps them.
pytestmark = pytest.mark.filterwarnings("ignore::UserWarning")

MAIN = {"m": producers.MAIN}
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"


def edit(book, sheet, ref, value):
    row, col = sf.parse_cell_name(ref)
    book.set_cells(sheet, {(row, col): book.blank_cell(sheet, row, col).edited(value=value)})


def work(book):
    """What a person does: types, inserts a row above the first data row,
    renames the last sheet."""
    first = book.sheets[0]
    edit(book, first, "B2", 1300)
    book.insert(first, "row", 1, 1)
    book.rename_sheet(book.sheets[-1], book.sheets[-1].name + " 2")
    book.save()


def part(path, name):
    with zipfile.ZipFile(path) as archive:
        return archive.read(name)


def sheet_root(path, number=1):
    return etree.fromstring(part(path, f"xl/worksheets/sheet{number}.xml"))


@pytest.mark.parametrize("producer", sorted(producers.PRODUCERS))
def test_an_ordinary_workbook_is_saved_into_its_own_file(tmp_path, producer):
    path = tmp_path / f"{producer}.xlsx"
    producers.PRODUCERS[producer](path)
    assert producers.check_package(path) == []
    assert sheet_io.unsaveable_parts(str(path)) == []
    book = load_workbook(str(path))
    work(book)
    assert producers.check_package(path) == []
    saved = openpyxl.load_workbook(path)
    first = saved.worksheets[0]
    assert first["B3"].value == 1300 and first["A3"].value == "Rent"
    assert first["B6"].value == "=SUM(B3:B5)"
    assert openpyxl.load_workbook(path, data_only=True).worksheets[0]["B6"].value == 1730
    assert saved.sheetnames[-1].endswith(" 2")
    assert sheet_io.unsaveable_parts(str(path)) == []
    # What MyEditor saved opens in MyEditor again, and saves again.
    again = load_workbook(str(path))
    edit(again, again.sheets[0], "A2", "Changed")
    again.save()
    assert producers.check_package(path) == []
    assert openpyxl.load_workbook(path).worksheets[0]["A2"].value == "Changed"


def test_excel_for_windows_keeps_and_moves_everything_it_wrote(tmp_path):
    path = tmp_path / "windows.xlsx"
    producers.excel_windows(path)
    book = load_workbook(str(path))
    work(book)                         # Lists becomes "Lists 2"
    root = sheet_root(path)
    # The error checks someone dismissed follow their cells.
    assert [(e.get("sqref"), sorted(e.attrib)) for e in root.iterfind(".//m:ignoredError", MAIN)] \
        == [("B3:B5", ["numberStoredAsText", "sqref"]), ("B6", ["formulaRange", "sqref"])]
    # The custom view's frozen pane and selection follow too.
    view = root.find(".//m:customSheetView", MAIN)
    assert view.find("m:pane", MAIN).get("topLeftCell") == "A4"
    assert view.find("m:selection", MAIN).get("sqref") == "B5"
    # Printer settings: the same bytes, linked from both page setups.
    assert part(path, "xl/printerSettings/printerSettings1.bin") == producers.PRINTER
    rels = etree.fromstring(part(path, "xl/worksheets/_rels/sheet1.xml.rels"))
    targets = {rel.get("Id"): rel.get("Target") for rel in rels}
    for setup in root.iter(f"{{{producers.MAIN}}}pageSetup"):
        assert targets[setup.get(f"{{{producers.R}}}id")].endswith(
            "printerSettings/printerSettings1.bin")
    # The data bar keeps its link to its newer part, which moved with it.
    rule = root.find(".//m:cfRule[@type='dataBar']", MAIN)
    assert rule.find(f".//{{{X14}}}id").text == producers.GUID_CF
    sqrefs = [e.text for e in root.iter(f"{{{XM}}}sqref")]
    formulas = [e.text for e in root.iter(f"{{{XM}}}f")]
    assert sqrefs == ["B3:B5", "D3:D5", "C3"]
    # Validation from the renamed sheet and the sparkline's data follow.
    assert formulas == ["'Lists 2'!$A$1:$A$2", "Data!B3:B5"]
    # The threaded comment moved with its cell; its person and its legacy
    # note are there.
    threaded = etree.fromstring(part(path, "xl/threadedComments/threadedComment1.xml"))
    assert [c.get("ref") for c in threaded] == ["B4"]
    assert openpyxl.load_workbook(path).worksheets[0]["B4"].comment.author == \
        f"tc={producers.GUID_THREAD}"
    book_rels = part(path, "xl/_rels/workbook.xml.rels").decode()
    assert "relationships/person" in book_rels
    assert b"Ann" in part(path, "xl/persons/person.xml")
    # The add-in: its parts byte for byte, linked from the package.
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
    assert part(path, "xl/webextensions/webextension1.xml").startswith(b"<?xml")
    assert "webextensiontaskpanes" in part(path, "_rels/.rels").decode()
    # The custom workbook view still names the sheet it showed.
    workbook = etree.fromstring(part(path, "xl/workbook.xml"))
    view = workbook.find(".//m:customWorkbookView", MAIN)
    sheet_ids = {s.get("name"): s.get("sheetId") for s in workbook.find("m:sheets", MAIN)}
    assert view.get("activeSheetId") == sheet_ids["Lists 2"]
    # The calculation chain is a cache Excel rebuilds; custom properties stay.
    assert "xl/calcChain.xml" not in names
    assert b"Budget" in part(path, "docProps/custom.xml")


def test_parts_kept_untouched_are_byte_for_byte_what_the_file_held(tmp_path):
    path = tmp_path / "windows.xlsx"
    producers.excel_windows(path)
    with zipfile.ZipFile(path) as archive:
        before = {name: archive.read(name) for name in archive.namelist()}
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "A2", "Rent!")
    book.save()
    for name in ("xl/printerSettings/printerSettings1.bin", "xl/persons/person.xml",
                 "xl/webextensions/webextension1.xml", "xl/webextensions/taskpanes.xml",
                 "xl/webextensions/_rels/taskpanes.xml.rels"):
        assert part(path, name) == before[name], name
    # openpyxl writes the custom properties again: the same property.
    assert openpyxl.load_workbook(path).custom_doc_props["Project"].value == "Budget"


def test_headers_and_footers_openpyxl_cannot_read_are_kept_as_written(tmp_path):
    for make, kind, text in (
            (producers.libreoffice, "oddHeader", '&C&"Times New Roman,Regular"&12&A'),
            (producers.numbers, "oddFooter", '&C&"Helvetica Neue,Regular"&12&K000000&P')):
        path = tmp_path / f"{make.__name__}.xlsx"
        make(path)
        book = load_workbook(str(path))
        work(book)
        assert sheet_root(path).find(f".//m:{kind}", MAIN).text == text


def test_undo_puts_the_kept_parts_back(tmp_path):
    path = tmp_path / "windows.xlsx"
    producers.excel_windows(path)
    book = load_workbook(str(path))
    snapshots = [book.insert(book.sheets[0], "row", 1, 2),
                 book.rename_sheet(book.sheets[1], "Choices")]
    for snapshot in reversed(snapshots):
        book.restore(snapshot)
    book.save()
    root = sheet_root(path)
    assert [e.get("sqref") for e in root.iterfind(".//m:ignoredError", MAIN)] == ["B2:B4", "B5"]
    assert [e.text for e in root.iter(f"{{{XM}}}f")] == ["Lists!$A$1:$A$2", "Data!B2:B4"]
    threaded = etree.fromstring(part(path, "xl/threadedComments/threadedComment1.xml"))
    assert [c.get("ref") for c in threaded] == ["B3"]


def test_deleting_the_cells_of_kept_rules_removes_those_rules_only(tmp_path):
    path = tmp_path / "windows.xlsx"
    producers.excel_windows(path)
    book = load_workbook(str(path))
    book.delete(book.sheets[0], "col", 2, 2)      # columns C and D: sparkline, validation
    book.save()
    assert producers.check_package(path) == []
    root = sheet_root(path)
    assert [e.text for e in root.iter(f"{{{XM}}}sqref")] == ["B2:B4"]
    assert root.find(f".//{{{X14}}}sparkline") is None
    assert root.find(f".//{{{X14}}}dataValidation") is None
    assert root.find(".//m:cfRule[@type='dataBar']", MAIN) is not None


def test_a_sheet_with_a_slicer_is_still_saved_only_as_a_copy(tmp_path):
    path = tmp_path / "slicer.xlsx"
    producers.excel_windows(path)
    data = part(path, "xl/worksheets/sheet1.xml").replace(
        b"</extLst>", b'<ext uri="{A8765BA9-456A-4dab-B4F3-ACF838C121DE}" '
                      b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
                      b"<x14:slicerList/></ext></extLst>")
    with zipfile.ZipFile(path) as archive:
        entries = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    entries["xl/worksheets/sheet1.xml"] = data
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    assert sheet_io.unsaveable_parts(str(path)) == ["Slicers"]


def test_the_element_order_check_catches_a_misplaced_element(tmp_path):
    path = tmp_path / "web.xlsx"
    producers.excel_web(path)
    with zipfile.ZipFile(path) as archive:
        entries = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    entries["xl/worksheets/sheet1.xml"] = re.sub(
        rb"(<sheetData>.*</sheetData>)(<pageMargins[^>]*/>)", rb"\2\1",
        entries["xl/worksheets/sheet1.xml"], flags=re.S)
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    assert any("out of order" in problem for problem in producers.check_package(path))
