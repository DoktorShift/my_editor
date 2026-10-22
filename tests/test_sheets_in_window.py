# SPDX-FileCopyrightText: 2026 rinbal
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Spreadsheets in the main window: the ones MyEditor edits open in a
sheet tab, which saves through the window (asking what each kind of save
needs), takes the Edit menu's commands and Format > Table and Format >
Sheet, marks its tab while it has unsaved changes, asks before it closes,
reads its file again when another app saves it, and comes back after an
update restart."""

from __future__ import annotations

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

import workspace  # noqa: E402
import workspace_restore  # noqa: E402
from tests.office_samples import make_zip  # noqa: E402
from tests.sheet_fixtures import realistic_ods, simple_xlsx  # noqa: E402
from tests.window_fixture import close_window, make_window  # noqa: E402


@pytest.fixture
def window(tmp_path, monkeypatch):
    QApplication.instance() or QApplication([])
    w = make_window(tmp_path, monkeypatch)
    yield w
    close_window(w)


@pytest.fixture
def xlsx(tmp_path):
    path = tmp_path / "Budget.xlsx"
    simple_xlsx(path)
    return str(path)


def test_the_spreadsheets_myeditor_edits_open_in_a_sheet_tab(window, tmp_path, xlsx):
    from sheet_tab import SheetTab
    from viewer_tabs import FileViewerTab
    ods = tmp_path / "Plan.ods"
    realistic_ods(ods)
    csv = tmp_path / "list.csv"
    csv.write_text("a,b\n1,2\n", encoding="utf-8")
    xls = make_zip(tmp_path / "Old.xls", {"x": "y"})
    numbers = make_zip(tmp_path / "Shop.numbers", {"Index/Document.iwa": "x"})
    for path in (xlsx, str(ods), str(csv)):
        assert isinstance(window.open_path(path), SheetTab), path
    for path in (xls, numbers):
        assert isinstance(window.open_path(path), FileViewerTab), path
    assert window.tabs.tabText(window.tabs.indexOf(window._sheet_tabs()[0])) == "Budget.xlsx"


def test_an_edit_marks_the_tab_and_save_clears_it(window, xlsx):
    sheet = window.open_path(xlsx)
    sheet.commit_text(0, 0, "42")
    index = window.tabs.indexOf(sheet)
    assert window.tabs.tabText(index) == "Budget.xlsx*"
    assert window.isWindowModified()
    assert window.save() is True
    assert window.tabs.tabText(index) == "Budget.xlsx" and not window.isWindowModified()
    assert openpyxl.load_workbook(xlsx).active["A1"].value == 42


def test_undo_and_redo_in_the_edit_menu_take_the_sheets_steps(window, xlsx):
    sheet = window.open_path(xlsx)
    before = sheet.model.cell(0, 0).value if sheet.model.cell(0, 0) else None
    sheet.commit_text(0, 0, "7")
    window._update_undo_redo_buttons()
    assert window.act_undo.isEnabled()
    window.act_undo.trigger()
    after = sheet.model.cell(0, 0)
    assert (after.value if after else None) == before
    window._update_undo_redo_buttons()
    assert window.act_redo.isEnabled()
    window.act_redo.trigger()
    assert sheet.model.cell(0, 0).value == 7


def test_copy_and_delete_act_on_the_cells(window, xlsx):
    sheet = window.open_path(xlsx)
    sheet.commit_text(0, 0, "hello")
    sheet.go_to(0, 0)
    window._update_editor_commands()
    assert window.act_select_all.isEnabled() and window.act_delete.isEnabled()
    window._edit_focused("copy")
    assert QApplication.clipboard().text().strip() == "hello"
    window._edit_focused("delete")
    assert sheet.model.cell(0, 0) is None or sheet.model.cell(0, 0).value in (None, "")
    QApplication.clipboard().clear()


def test_format_table_and_sheet_commands_act_on_the_workbook(window, xlsx):
    sheet = window.open_path(xlsx)
    window._update_editor_commands()
    add_row = window.commands.action("format.table.add_row_above")
    assert add_row.isEnabled()
    assert not window.commands.action("format.table.delete_table").isEnabled()
    rows_before = sheet.sheet.max_row if hasattr(sheet.sheet, "max_row") else None
    sheet.commit_text(0, 0, "top")
    sheet.go_to(0, 0)
    add_row.trigger()
    assert sheet.model.cell(1, 0).value == "top"
    sheets = len(sheet.book.sheets)
    window.commands.action("format.sheet.add").trigger()
    assert len(sheet.book.sheets) == sheets + 1
    del rows_before


def test_sheet_commands_are_off_where_there_is_no_sheet(window):
    window.new_tab()
    window._update_editor_commands()
    assert not window.commands.action("format.sheet.add").isEnabled()


@pytest.mark.parametrize("status, answer, expected", [
    ("in_use", None, ["inform"]),
    ("failed", None, ["inform"]),
    ("changed_on_disk", "replace", ["ask", "save(overwrite)"]),
    ("changed_on_disk", "revert", ["ask", "reload"]),
    ("needs_copy", "copy", ["ask", "save_as"]),
    ("needs_copy", "cancel", ["ask"]),
    ("deleted", "save_as", ["ask", "save_as"]),
    ("deleted", "cancel", ["ask"]),
])
def test_each_way_a_save_can_go_is_answered(window, xlsx, monkeypatch, status, answer,
                                            expected):
    import sheet_tab
    import sheet_window
    sheet = window.open_path(xlsx)
    done = []
    monkeypatch.setattr(sheet_window, "inform", lambda *a, **k: done.append("inform"))

    def ask(*_a, buttons, **_k):
        done.append("ask")
        if status == "needs_copy":
            labels = [b.label for b in buttons]
            assert labels[-1] == "Save a Copy…"         # the default, last
        return answer

    monkeypatch.setattr(sheet_window, "ask", ask)
    results = iter([sheet_tab.SaveResult(status, "why"), sheet_tab.SaveResult(sheet_tab.SAVED)])
    monkeypatch.setattr(sheet, "save", lambda overwrite_changes=False: (
        done.append("save(overwrite)") if overwrite_changes else None) or next(results))
    monkeypatch.setattr(sheet, "reload", lambda: done.append("reload"))
    monkeypatch.setattr(window._sheet_window, "save_as", lambda s: done.append("save_as") or True)
    window.save()
    assert done == expected


def test_closing_a_changed_sheet_asks_first(window, xlsx, monkeypatch):
    import main_window
    sheet = window.open_path(xlsx)
    sheet.commit_text(0, 0, "12345")
    assert sheet.is_modified()
    monkeypatch.setattr(main_window, "_ask_save_changes", lambda *a, **k: "cancel")
    window.close_tab(window.tabs.indexOf(sheet))
    assert window.tabs.indexOf(sheet) >= 0
    monkeypatch.setattr(main_window, "_ask_save_changes", lambda *a, **k: "discard")
    window.close_tab(window.tabs.indexOf(sheet))
    assert window.tabs.indexOf(sheet) < 0


def test_a_sheet_hears_when_another_app_saves_its_file(window, xlsx, monkeypatch):
    sheet = window.open_path(xlsx)
    heard = []
    monkeypatch.setattr(sheet, "schedule_reload", lambda: heard.append(True))
    window._on_file_changed(xlsx)
    assert heard == [True]


def test_a_sheet_comes_back_after_an_update_and_a_changed_one_is_asked_about(window, xlsx,
                                                                            tmp_path):
    sheet = window.open_path(xlsx)
    capture = workspace_restore.capture_tabs(window)
    assert [t.kind for t in capture.tabs()[0]][-1] == workspace.SHEET
    sheet.commit_text(0, 0, "3")
    capture = workspace_restore.capture_tabs(window)
    assert capture.unprotected == [sheet]
    capture.settle(sheet, saved=False)
    tabs, active = capture.tabs()
    assert tabs[-1].kind == workspace.SHEET and tabs[-1].path == sheet._file_path
    record = tmp_path / "ws.json"
    workspace.write_workspace(workspace.Workspace(tabs=tabs, active=active), str(record))
    sheet.undo_stack.setClean()
    window.close_tab(window.tabs.indexOf(sheet))
    workspace_restore.resume(window, workspace.take_workspace(str(record)), draft_type=dict)
    assert window._sheet_tabs() and window._sheet_tabs()[0]._file_path == sheet._file_path


def test_renaming_a_sheets_file_keeps_its_changes(window, xlsx, tmp_path, monkeypatch):
    import main_window
    sheet = window.open_path(xlsx)
    sheet.commit_text(0, 0, "9")
    monkeypatch.setattr(main_window.QInputDialog, "getText",
                        staticmethod(lambda *a, **k: ("Costs.xlsx", True)))
    window._rename_viewer_tab(window.tabs.indexOf(sheet), sheet)
    assert sheet._file_path == str(tmp_path / "Costs.xlsx") and sheet.is_modified()
    assert window.tabs.tabText(window.tabs.indexOf(sheet)) == "Costs.xlsx*"
    assert window.save()
    assert openpyxl.load_workbook(tmp_path / "Costs.xlsx").active["A1"].value == 9


def test_formatting_and_inserting_text_are_off_in_a_sheet(window, xlsx):
    window.open_path(xlsx)
    window._update_editor_commands()
    enabled = sorted(c.id for c in window.commands.commands()
                     if c.id.startswith(("format.", "insert."))
                     and not c.id.startswith(("format.table.", "format.sheet."))
                     and window.commands.action(c.id).isEnabled())
    assert enabled == []


def test_saving_before_an_update_saves_a_changed_sheet_too(window, xlsx, monkeypatch):
    import main_window
    sheet = window.open_path(xlsx)
    sheet.commit_text(0, 0, "777")
    note = window.new_tab()
    note.insertPlainText("words")
    monkeypatch.setattr(main_window, "_ask_save_changes", lambda *a, **k: "save")
    monkeypatch.setattr(window, "save_as", lambda *a, **k: (
        note.document().setModified(False) or True))
    assert window._confirm_save_before_update()
    assert openpyxl.load_workbook(xlsx).active["A1"].value == 777
    assert window._unsaved_tabs() == []


def test_a_missing_spreadsheet_is_said_once(window, tmp_path, monkeypatch):
    import main_window
    said = []
    monkeypatch.setattr(main_window, "inform", lambda *a, **k: said.append(k.get("title")))
    assert window.open_path(str(tmp_path / "Gone.xlsx")) is None
    assert len(said) == 1


def test_save_as_offers_what_the_workbook_can_be_saved_as_and_follows_the_pick(
        window, xlsx, tmp_path, monkeypatch):
    import main_window
    sheet = window.open_path(xlsx)
    asked = {}

    def dialog(parent, title, suggested, filters, selected):
        asked["filters"] = filters.split(";;")
        asked["suggested"] = suggested
        # The person keeps the suggested name and picks the CSV format.
        return suggested, asked["filters"][1]

    monkeypatch.setattr(main_window.QFileDialog, "getSaveFileName", dialog)
    assert window._sheet_window.save_as(sheet)
    assert asked["filters"] == ["Excel Workbook (*.xlsx)", "CSV, the Sheet Shown (*.csv)",
                                "TSV, the Sheet Shown (*.tsv)"]
    assert asked["suggested"].endswith("Budget 2.xlsx")
    assert (tmp_path / "Budget 2.csv").exists()
    assert not (tmp_path / "Budget 2.xlsx").exists()
    assert sheet._file_path == xlsx


def test_a_sheet_whose_file_went_to_the_trash_offers_save_as_and_close(window, xlsx,
                                                                      monkeypatch):
    sheet = window.open_path(xlsx)
    sheet.commit_text(0, 0, "42")
    os.remove(xlsx)
    sheet._do_scheduled_reload()
    assert sheet.is_deleted()
    assert not sheet._save_as_button.isHidden() and sheet._reload_button.isHidden()
    asked = []
    monkeypatch.setattr(window._sheet_window, "save_as", lambda s: asked.append(s) or True)
    sheet._save_as_button.click()
    assert asked == [sheet]
    monkeypatch.setattr("main_window.ask", lambda *a, **k: "discard")
    sheet._close_button.click()
    assert sheet not in window._sheet_tabs()


def test_a_sheet_closed_while_it_saves_closes_once_saved(window, xlsx, monkeypatch):
    import threading

    import shiboken6
    from PySide6.QtCore import QTimer
    from PySide6.QtTest import QTest

    import sheet_tab
    monkeypatch.setattr(sheet_tab, "BACKGROUND_CELLS", -1)
    sheet = window.open_path(xlsx)
    sheet.commit_text(0, 0, "42")
    release = threading.Event()
    save = sheet.book.save
    monkeypatch.setattr(sheet.book, "save", lambda *a, **k: release.wait(5) and save(*a, **k))
    asked, closed_at_once = [], []
    monkeypatch.setattr("main_window.ask", lambda *a, **k: asked.append(k) or "cancel")

    def close_now():
        window.close_tab(window.tabs.indexOf(sheet))
        closed_at_once.append(window.tabs.indexOf(sheet) < 0)
        release.set()

    timer = QTimer()
    timer.setSingleShot(True)
    timer.timeout.connect(close_now)
    timer.start(50)
    assert window.save()
    assert closed_at_once == [False]
    timer.timeout.disconnect(close_now)
    def gone():
        return not shiboken6.isValid(sheet) or window.tabs.indexOf(sheet) < 0

    for _ in range(100):
        if gone():
            break
        QTest.qWait(20)
    assert gone()
    assert openpyxl.load_workbook(xlsx).active["A1"].value == 42
    assert asked == []
