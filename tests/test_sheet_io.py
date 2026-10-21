# SPDX-FileCopyrightText: 2026 rinbal
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Spreadsheet files keep everything that is in them when MyEditor saves.

These tests guard people's data: what a save must keep (styles, merged
cells, names, macros, every other part of an OpenDocument package, a text
table's delimiter, encoding and quotes), and when it must not write at
all (another app has the file open, the file changed on the disk, a part
would be lost).
"""

import datetime as dt
import os
import re
import zipfile

import openpyxl
import pytest
from lxml import etree
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import sheet_formulas as sf
import sheet_io
from sheet_io import NumberStyle, format_value, load_workbook
from tests.sheet_fixtures import ODS_MIMETYPE, minimal_ods, realistic_ods, simple_xlsx

TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"


def edit(book, sheet, ref, value=None, formula=None, number_format=None):
    row, col = sf.parse_cell_name(ref)
    cell = book.blank_cell(sheet, row, col)
    new = cell.edited(value=value, formula=formula)
    if number_format:
        new = book.cell_with_format(sheet, new, number_format)
    return book.set_cells(sheet, {(row, col): new})


def value(sheet, ref):
    cell = sheet.cells.get(sf.parse_cell_name(ref))
    return None if cell is None else cell.value


# --------------------------------------------------------------------------- #
# Excel                                                                       #
# --------------------------------------------------------------------------- #

@pytest.fixture
def styled_xlsx(tmp_path):
    path = tmp_path / "Budget.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Item"
    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["B1"] = "Amount"
    ws["A2"], ws["B2"] = "Rent", 1200
    ws["A3"], ws["B3"] = "Food", 350.5
    ws["B2"].number_format = "#,##0.00 [$€-407]"
    ws["B3"].number_format = "#,##0.00 [$€-407]"
    ws["B3"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["A4"], ws["B4"] = "Total", "=SUM(B2:B3)"
    ws["C4"] = "=B3/B4"
    ws["C4"].number_format = "0%"
    ws["D1"] = dt.datetime(2026, 10, 21)
    ws["D1"].number_format = "yyyy-mm-dd"
    ws.merge_cells("E1:F2")
    ws["E1"] = "Merged"
    ws["A2"].comment = Comment("Due on the first", "Ann")
    ws.column_dimensions["A"].width = 24
    ws.freeze_panes = "A2"
    validation = DataValidation(type="whole", operator="greaterThan", formula1="0")
    validation.add("B2:B3")
    ws.add_data_validation(validation)
    ws.conditional_formatting.add("B2:B3", CellIsRule(operator="greaterThan", formula=["1000"],
                                                      fill=PatternFill("solid",
                                                                       fgColor="00FF00")))
    wb.defined_names["Amounts"] = DefinedName("Amounts", attr_text="Data!$B$2:$B$3")
    other = wb.create_sheet("Summary")
    other["A1"] = "=SUM(Data!B2:B3)*2"
    other["A2"] = "=SUM(Amounts)"
    wb.save(path)
    return str(path)


def test_an_edited_workbook_keeps_its_styles_merges_formats_and_names(styled_xlsx):
    book = load_workbook(styled_xlsx)
    data = book.sheets[0]
    edit(book, data, "B2", 1300)
    book.save()

    wb = openpyxl.load_workbook(styled_xlsx)
    ws = wb["Data"]
    assert ws["B2"].value == 1300
    assert ws["A1"].font.b and ws["A1"].font.color.rgb.endswith("FF0000")
    assert ws["B3"].fill.fgColor.rgb.endswith("FFFF00")
    assert ws["B2"].number_format == "#,##0.00 [$€-407]"
    assert ws["D1"].value == dt.datetime(2026, 10, 21)
    assert [str(r) for r in ws.merged_cells.ranges] == ["E1:F2"]
    assert ws["A2"].comment.text == "Due on the first"
    assert ws.column_dimensions["A"].width == 24
    assert ws.freeze_panes == "A2"
    assert str(ws.data_validations.dataValidation[0].sqref) == "B2:B3"
    assert [str(cf.sqref) for cf in ws.conditional_formatting] == ["B2:B3"]
    assert wb.defined_names["Amounts"].attr_text == "Data!$B$2:$B$3"
    assert ws["B4"].value == "=SUM(B2:B3)"


def test_formula_results_are_written_for_other_apps_and_previews(styled_xlsx):
    book = load_workbook(styled_xlsx)
    assert value(book.sheets[0], "B4") == 1550.5       # calculated: openpyxl stored none
    edit(book, book.sheets[0], "B2", 1300)
    assert value(book.sheets[0], "B4") == 1650.5
    assert value(book.sheets[1], "A1") == 3301
    book.save()
    cached = openpyxl.load_workbook(styled_xlsx, data_only=True)
    assert cached["Data"]["B4"].value == 1650.5
    assert cached["Data"]["C4"].value == pytest.approx(350.5 / 1650.5)
    assert cached["Summary"]["A1"].value == 3301
    assert cached["Summary"]["A2"].value == 1650.5
    with zipfile.ZipFile(styled_xlsx) as archive:
        workbook_xml = archive.read("xl/workbook.xml")
    assert b'fullCalcOnLoad="1"' in workbook_xml


def test_text_and_error_results_carry_their_type(tmp_path):
    path = tmp_path / "types.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = '="a"&"b"'
    wb.active["A2"] = "=1/0"
    wb.active["A3"] = "=1>0"
    wb.save(path)
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "B1", 1)
    book.save()
    cached = openpyxl.load_workbook(path, data_only=True).active
    assert cached["A1"].value == "ab"
    assert cached["A2"].value == "#DIV/0!"
    assert cached["A3"].value is True


def test_a_function_myeditor_does_not_know_keeps_its_stored_result(tmp_path):
    path = tmp_path / "unknown.xlsx"
    simple_xlsx(path)
    with zipfile.ZipFile(path) as archive:
        entries = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    sheet_xml = entries["xl/worksheets/sheet1.xml"].replace(
        b'<c r="A3"><f>A1+A2</f><v></v></c>',
        b'<c r="A3"><f>_xlfn.FORECAST.ETS(A1,A1:A2,A1:A2)</f><v>7.5</v></c>')
    assert sheet_xml != entries["xl/worksheets/sheet1.xml"]
    entries["xl/worksheets/sheet1.xml"] = sheet_xml
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    book = load_workbook(str(path))
    cell = book.sheets[0].cells[(2, 0)]
    assert cell.value == 7.5 and cell.state == sf.STATE_APP
    edit(book, book.sheets[0], "B1", "x")
    book.save()
    assert openpyxl.load_workbook(path, data_only=True).active["A3"].value == 7.5


def test_inserting_a_row_moves_formulas_merges_names_and_validation(styled_xlsx):
    book = load_workbook(styled_xlsx)
    data = book.sheets[0]
    book.insert(data, "row", 1, 1)
    assert data.cells[(4, 1)].formula == "=SUM(B3:B4)"
    assert book.sheets[1].cells[(0, 0)].formula == "=SUM(Data!B3:B4)*2"
    book.save()
    wb = openpyxl.load_workbook(styled_xlsx)
    ws = wb["Data"]
    assert ws["A3"].value == "Rent" and ws["A3"].comment.text == "Due on the first"
    assert [str(r) for r in ws.merged_cells.ranges] == ["E1:F3"]
    assert wb.defined_names["Amounts"].attr_text == "Data!$B$3:$B$4"
    assert str(ws.data_validations.dataValidation[0].sqref) == "B3:B4"
    assert [str(cf.sqref) for cf in ws.conditional_formatting] == ["B3:B4"]
    # Inserted at the first scrolling row: the frozen row stays one row, as in Excel.
    assert ws.freeze_panes == "A2"
    assert openpyxl.load_workbook(styled_xlsx, data_only=True)["Data"]["B5"].value == 1550.5


def test_deleting_a_row_that_a_formula_reads_shows_ref(styled_xlsx):
    book = load_workbook(styled_xlsx)
    data = book.sheets[0]
    book.insert(data, "col", 0, 1)
    book.delete(data, "row", 2, 1)
    assert data.cells[(2, 3)].formula == "=#REF!/C3"
    book.delete(data, "col", 2, 1)
    assert data.cells[(2, 2)].formula == "=#REF!/#REF!"
    assert data.cells[(2, 2)].value == sf.REF


def test_undo_of_a_structural_change_puts_everything_back(styled_xlsx):
    book = load_workbook(styled_xlsx)
    data = book.sheets[0]
    before = dict(data.cells)
    snapshot = book.delete(data, "row", 1, 2)
    assert value(book.sheets[1], "A1") == sf.REF
    book.restore(snapshot)
    assert data.cells == before
    assert value(book.sheets[1], "A1") == 3101
    book.save()
    wb = openpyxl.load_workbook(styled_xlsx)
    assert wb.defined_names["Amounts"].attr_text == "Data!$B$2:$B$3"
    assert [str(r) for r in wb["Data"].merged_cells.ranges] == ["E1:F2"]


def test_sheets_are_added_renamed_duplicated_and_deleted_in_excel(styled_xlsx):
    book = load_workbook(styled_xlsx)
    book.add_sheet("Notes")
    book.rename_sheet(book.sheets[0], "Money")
    assert book.sheets[1].cells[(0, 0)].formula == "=SUM(Money!B2:B3)*2"
    book.duplicate_sheet(book.sheets[0])
    snapshot = book.delete_sheet(book.sheets[1])
    assert [s.name for s in book.sheets] == ["Money", "Summary", "Notes"]
    book.restore(snapshot)
    assert [s.name for s in book.sheets] == ["Money", "Money 2", "Summary", "Notes"]
    book.save()
    wb = openpyxl.load_workbook(styled_xlsx)
    assert wb.sheetnames == ["Money", "Money 2", "Summary", "Notes"]
    assert wb["Summary"]["A1"].value == "=SUM(Money!B2:B3)*2"
    assert wb.defined_names["Amounts"].attr_text == "Money!$B$2:$B$3"
    assert wb["Money 2"]["A1"].font.b


def test_a_sheet_name_is_checked_as_excel_checks_it(styled_xlsx):
    book = load_workbook(styled_xlsx)
    assert book.sheet_name_problem("summary") is not None
    assert book.sheet_name_problem("a/b") is not None
    assert book.sheet_name_problem("x" * 32) is not None
    assert book.sheet_name_problem("Fine") is None
    with pytest.raises(sheet_io.SheetFileError):
        book.rename_sheet(book.sheets[0], "Summary")


def test_column_widths_are_saved_into_the_file(styled_xlsx):
    book = load_workbook(styled_xlsx)
    book.set_width(book.sheets[0], 1, 150)
    book.save()
    ws = openpyxl.load_workbook(styled_xlsx)["Data"]
    assert ws.column_dimensions["B"].width == pytest.approx((150 - 5) / 7, abs=0.01)
    assert ws.column_dimensions["A"].width == 24


def test_a_macro_workbook_keeps_its_macros(tmp_path):
    plain = tmp_path / "plain.xlsx"
    simple_xlsx(plain)
    path = tmp_path / "Macros.xlsm"
    vba = b"\xd0\xcf\x11\xe0 pretend VBA project" * 20
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(path, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml").replace(
                    b"</Types>", b'<Default Extension="bin" ContentType="application/'
                                 b'vnd.ms-office.vbaProject"/></Types>')
            if info.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>",
                                    b'<Relationship Id="rIdVBA" Type="http://schemas.microsoft'
                                    b'.com/office/2006/relationships/vbaProject" '
                                    b'Target="vbaProject.bin"/></Relationships>')
            target.writestr(info, data)
        target.writestr("xl/vbaProject.bin", vba)
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "C1", 5)
    book.save()
    with zipfile.ZipFile(path) as archive:
        assert archive.read("xl/vbaProject.bin") == vba
        assert b"macroEnabled" in archive.read("[Content_Types].xml")
    book.close()


def _with_parts(source, target, extra):
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(target, "w") as out:
        for info in archive.infolist():
            out.writestr(info, archive.read(info.filename))
        for name, data in extra.items():
            out.writestr(name, data)


def test_parts_openpyxl_would_drop_are_named(tmp_path):
    plain = tmp_path / "plain.xlsx"
    simple_xlsx(plain)
    assert sheet_io.unsaveable_parts(str(plain)) == []
    rich = tmp_path / "rich.xlsx"
    _with_parts(plain, rich, {
        "xl/charts/chart1.xml": b"<c/>",
        "xl/drawings/drawing1.xml": b"<xdr:wsDr><xdr:pic/></xdr:wsDr>",
        "xl/media/image1.png": b"png",
        "xl/pivotTables/pivotTable1.xml": b"<p/>",
        "xl/slicers/slicer1.xml": b"<s/>",
    })
    parts = sheet_io.unsaveable_parts(str(rich))
    assert parts[:2] == ["Charts", "Pictures"]
    assert "Pivot tables" in parts and "Slicers" in parts


def test_a_workbook_with_parts_that_would_be_lost_is_never_saved_over(tmp_path):
    plain = tmp_path / "plain.xlsx"
    simple_xlsx(plain)
    path = tmp_path / "Report.xlsx"
    _with_parts(plain, path, {"xl/charts/chart1.xml": b"<c/>"})
    before = path.read_bytes()
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "A1", 100)
    with pytest.raises(sheet_io.UnsaveablePartsError) as raised:
        book.save()
    assert raised.value.parts == ["Charts"]
    assert path.read_bytes() == before
    copy = tmp_path / "Report copy.xlsx"
    book.save(str(copy))
    assert openpyxl.load_workbook(copy).active["A1"].value == 100
    book.adopt_path(str(copy))
    book.save()


def test_a_damaged_file_gives_a_sentence_not_a_crash(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a zip")
    with pytest.raises(sheet_io.SheetFileError):
        load_workbook(str(path))


@pytest.mark.skipif(os.name == "nt", reason="named pipes")
@pytest.mark.parametrize("name", ["Budget.csv", "Budget.xlsx", "Budget.ods"])
def test_a_pipe_with_a_spreadsheets_name_is_never_opened(tmp_path, name):
    import threading
    os.mkfifo(tmp_path / name)
    outcome = []

    def load():
        try:
            load_workbook(str(tmp_path / name))
        except sheet_io.SheetFileError as error:
            outcome.append(str(error))
        outcome.append(sheet_io.unsaveable_parts(str(tmp_path / name)))

    thread = threading.Thread(target=load, daemon=True)
    thread.start()
    thread.join(3)
    assert not thread.is_alive()
    assert outcome == ["The file is damaged or is not a spreadsheet.", []]


# --------------------------------------------------------------------------- #
# The disk                                                                  #
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("lock_name", ["~$Budget.xlsx", ".~lock.Budget.xlsx#"])
def test_a_lock_file_of_another_app_blocks_writing(styled_xlsx, lock_name):
    folder = os.path.dirname(styled_xlsx)
    book = load_workbook(styled_xlsx)
    edit(book, book.sheets[0], "B2", 1)
    before = open(styled_xlsx, "rb").read()
    lock = os.path.join(folder, lock_name)
    open(lock, "w", encoding="utf-8").close()
    assert book.lock_owner() == lock
    with pytest.raises(sheet_io.FileInUseError):
        book.save()
    assert open(styled_xlsx, "rb").read() == before
    os.remove(lock)
    book.save()


def test_excel_shortens_the_lock_name_of_a_long_file_name(tmp_path):
    path = tmp_path / "Quarterly.xlsx"
    simple_xlsx(path)
    (tmp_path / "~$arterly.xlsx").write_text("owner", encoding="utf-8")
    book = load_workbook(str(path))
    assert book.lock_owner() == str(tmp_path / "~$arterly.xlsx")
    with pytest.raises(sheet_io.FileInUseError):
        book.save()


def test_a_file_that_changed_on_the_disk_is_not_overwritten_unasked(styled_xlsx):
    book = load_workbook(styled_xlsx)
    assert not book.changed_on_disk()
    edit(book, book.sheets[0], "B2", 1)
    wb = openpyxl.load_workbook(styled_xlsx)
    wb["Data"]["A6"] = "written elsewhere"
    wb.save(styled_xlsx)
    os.utime(styled_xlsx, ns=(1, 1))
    assert book.changed_on_disk()
    with pytest.raises(sheet_io.ChangedOnDiskError):
        book.save()
    assert openpyxl.load_workbook(styled_xlsx)["Data"]["A6"].value == "written elsewhere"
    book.save(overwrite_changes=True)
    assert not book.changed_on_disk()


def test_saving_keeps_the_files_permissions(styled_xlsx):
    os.chmod(styled_xlsx, 0o640)
    book = load_workbook(styled_xlsx)
    edit(book, book.sheets[0], "B2", 1)
    book.save()
    assert os.stat(styled_xlsx).st_mode & 0o777 == 0o640


# --------------------------------------------------------------------------- #
# OpenDocument                                                                #
# --------------------------------------------------------------------------- #

def _entries(path):
    with zipfile.ZipFile(path) as archive:
        return [(i.filename, i.compress_type, archive.read(i.filename))
                for i in archive.infolist()]


def _table(path, name):
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("content.xml"))
    for table in root.iter(f"{{{TABLE}}}table"):
        if table.get(f"{{{TABLE}}}name") == name:
            return table
    raise KeyError(name)


def test_an_ods_reads_values_formulas_formats_and_layout(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    budget, notes = book.sheets
    assert [s.name for s in book.sheets] == ["Budget", "Notes"]
    assert value(budget, "B2") == 1200
    assert budget.cells[(1, 1)].shown == "1.200,00 €"
    assert sheet_io.number_format_kind(budget.cells[(1, 1)].number_format) == "currency"
    assert budget.cells[(6, 1)].formula == "=SUM(B2:B5)"
    assert budget.cells[(6, 2)].number_format == "0%"
    assert notes.cells[(0, 0)].formula == "=Budget!B7*2"
    assert budget.frozen == (1, 0)
    assert budget.widths[0] == pytest.approx(5 / 2.54 * 96)
    assert book.backend.names["Amounts"] == "=Budget!$B$2:$B$5"


def test_an_ods_round_trip_keeps_every_other_entry_byte_for_byte(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    before = _entries(path)
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "B3", 400)
    book.save()
    after = _entries(path)
    assert after[0][0] == "mimetype" and after[0][1] == zipfile.ZIP_STORED
    assert after[0][2] == ODS_MIMETYPE
    assert [(n, c) for n, c, _d in after] == [(n, c) for n, c, _d in before]
    for (name, _c, data), (_n, _c2, old) in zip(after, before):
        if name != "content.xml":
            assert data == old, name
    reread = load_workbook(str(path))
    budget = reread.sheets[0]
    assert value(budget, "B3") == 400
    assert value(budget, "B7") == 1600
    assert reread.sheets[1].cells[(0, 0)].value == 3200
    # The comment on A2 and the other cells' styles stay.
    assert b"Check this" in dict((n, d) for n, _c, d in after)["content.xml"]
    assert budget.cells[(1, 1)].style == "ce2"


def test_an_unchanged_ods_sheet_is_left_as_it_was(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    original = etree.tostring(_table(path, "Notes"))
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "A1", "Things")
    book.sheets[1].changed = False
    book.save()
    assert etree.tostring(_table(path, "Notes")) == original


def test_an_edit_inside_repeated_rows_and_cells_splits_only_that_run(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    budget = book.sheets[0]
    edit(book, budget, "B5", 99.5)          # the middle row of three repeated rows
    edit(book, budget, "D9", "far away")    # inside the last run of a million rows
    book.save()
    table = _table(path, "Budget")
    rows = []
    for row in table.iter(f"{{{TABLE}}}table-row"):
        rows.append(int(row.get(f"{{{TABLE}}}number-rows-repeated", "1")))
    assert sum(rows) == 1048576
    reread = load_workbook(str(path))
    sheet = reread.sheets[0]
    assert value(sheet, "B5") == 99.5 and value(sheet, "D9") == "far away"
    # The yellow cells around the edit keep their style.
    backend = reread.backend
    assert backend._style_at(sheet.backend, 3, 1) == "ce4"
    assert backend._style_at(sheet.backend, 4, 1) == "ce4"
    assert backend._style_at(sheet.backend, 5, 0) == "ce4"
    assert value(sheet, "B7") == 1650


def test_inserting_and_deleting_rows_and_columns_in_an_ods(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    budget = book.sheets[0]
    book.insert(budget, "row", 1, 2)
    book.insert(budget, "col", 0, 1)
    assert budget.cells[(8, 2)].formula == "=SUM(C4:C7)"
    assert book.sheets[1].cells[(0, 0)].formula == "=Budget!C9*2"
    book.save()
    reread = load_workbook(str(path))
    sheet = reread.sheets[0]
    assert value(sheet, "C4") == 1200 and value(sheet, "B4") == "Rent"
    assert value(sheet, "A4") is None
    assert reread.backend.names["Amounts"] == "=Budget!$C$4:$C$7"
    table = _table(path, "Budget")
    total = sum(int(r.get(f"{{{TABLE}}}number-rows-repeated", "1"))
                for r in table.iter(f"{{{TABLE}}}table-row"))
    assert total == 1048576
    snapshot = reread.delete(sheet, "row", 0, 3)
    assert value(sheet, "B1") == "Rent"
    reread.restore(snapshot)
    reread.delete(sheet, "col", 0, 1)
    reread.save()
    third = load_workbook(str(path))
    assert value(third.sheets[0], "B4") == 1200
    assert third.sheets[1].cells[(0, 0)].formula == "=Budget!B9*2"


def test_ods_sheets_are_added_renamed_and_deleted(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    book.add_sheet("Ideas")
    edit(book, book.sheets[2], "A1", "first idea")
    book.rename_sheet(book.sheets[0], "My Budget")
    book.duplicate_sheet(book.sheets[1])
    book.save()
    reread = load_workbook(str(path))
    assert [s.name for s in reread.sheets] == ["My Budget", "Notes", "Notes 2", "Ideas"]
    assert reread.sheets[1].cells[(0, 0)].formula == "='My Budget'!B7*2"
    assert value(reread.sheets[3], "A1") == "first idea"
    assert reread.backend.names["Amounts"] == "='My Budget'!$B$2:$B$5"
    snapshot = reread.delete_sheet(reread.sheets[0])
    assert reread.sheets[0].cells[(0, 0)].formula == "=#REF!*2"
    reread.restore(snapshot)
    reread.save()
    assert load_workbook(str(path)).sheets[1].cells[(0, 0)].value == 3101


def test_an_undo_after_a_save_still_writes_the_truth(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    budget = book.sheets[0]
    original = budget.cells[(0, 0)]
    edit(book, budget, "A1", "Changed")
    book.save()
    book.set_cells(budget, {(0, 0): original})    # what Undo does
    book.save()
    assert value(load_workbook(str(path)).sheets[0], "A1") == "Item"


def test_a_percent_typed_into_an_ods_gets_a_percent_style(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "D2", 0.15, number_format="0%")
    book.save()
    reread = load_workbook(str(path))
    cell = reread.sheets[0].cells[(1, 3)]
    assert cell.value == 0.15 and cell.number_format == "0%"
    assert cell.shown == "15%"


def test_a_minimal_ods_with_repeated_columns_is_edited(tmp_path):
    path = tmp_path / "plain.ods"
    minimal_ods(path, '<table:table-row table:number-rows-repeated="2">'
                      '<table:table-cell office:value-type="float" office:value="1" '
                      'table:number-columns-repeated="3"><text:p>1</text:p></table:table-cell>'
                      '</table:table-row>')
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    assert len(sheet.cells) == 6
    edit(book, sheet, "B2", None, formula="=SUM(A1:C1)")
    book.save()
    reread = load_workbook(str(path))
    assert value(reread.sheets[0], "B2") == 3
    assert value(reread.sheets[0], "A2") == 1 and value(reread.sheets[0], "C2") == 1
    assert value(reread.sheets[0], "B1") == 1


def test_an_ods_saved_as_xlsx_keeps_values_and_formulas(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    target = tmp_path / "Budget.xlsx"
    book.save(str(target))
    wb = openpyxl.load_workbook(target)
    assert wb.sheetnames == ["Budget", "Notes"]
    assert wb["Budget"]["B7"].value == "=SUM(B2:B5)"
    assert openpyxl.load_workbook(target, data_only=True)["Budget"]["B7"].value == 1550.5


# --------------------------------------------------------------------------- #
# Text tables                                                                 #
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("data", [
    b"a,b,c\r\n1,2,3\r\n",
    b"\xef\xbb\xbfName;Wert\n\"M\xc3\xbcller, Hans\";\"12,5\"\n;\n",
    b"x\ty\n\"quoted\"\t\"with \"\"quotes\"\"\"\n",
    b'"a","b"\n"1",""\n',
    b"one,two\nthree",
    b"a,b\n\n,\nc\n",
])
def test_an_unchanged_text_table_saves_byte_for_byte(tmp_path, data):
    path = tmp_path / "table.csv"
    path.write_bytes(data)
    book = load_workbook(str(path))
    book.sheets[0].changed = True
    book.save()
    assert path.read_bytes() == data


def test_a_text_table_keeps_its_delimiter_encoding_and_quotes(tmp_path):
    path = tmp_path / "Liste.csv"
    path.write_bytes(b"\xef\xbb\xbfName;Wert\r\n\"M\xc3\xbcller\";12,5\r\n")
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    assert value(sheet, "B2") == "12,5"
    edit(book, sheet, "A2", "Schmidt; Anna")
    edit(book, sheet, "B3", "=1+1")
    book.save()
    assert path.read_bytes() == ("﻿Name;Wert\r\n\"Schmidt; Anna\";12,5\r\n;=1+1\r\n"
                                 .encode("utf-8"))
    assert value(load_workbook(str(path)).sheets[0], "B3") == "=1+1"


def test_a_windows_text_table_keeps_its_encoding(tmp_path):
    path = tmp_path / "old.csv"
    path.write_bytes("Größe,Preis\n".encode("cp1252"))
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "A2", "ä")
    book.save()
    assert path.read_bytes() == "Größe,Preis\nä,\n".encode("cp1252")
    edit(book, book.sheets[0], "A3", "中")
    with pytest.raises(sheet_io.SheetFileError):
        book.save()


def test_rows_and_columns_in_a_text_table(tmp_path):
    path = tmp_path / "t.tsv"
    path.write_bytes(b"a\tb\tc\n1\t2\t3\n")
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    book.insert(sheet, "col", 1, 1)
    book.insert(sheet, "row", 0, 1)
    book.save()
    assert path.read_bytes() == b"\t\t\t\na\t\tb\tc\n1\t\t2\t3\n"
    with pytest.raises(sheet_io.SheetFileError):
        book.add_sheet()


def test_a_workbook_exports_its_sheet_as_csv(styled_xlsx, tmp_path):
    book = load_workbook(styled_xlsx)
    target = tmp_path / "out.csv"
    book.save(str(target), sheet=book.sheets[0])
    text = target.read_bytes().decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Item,Amount,")
    assert "Total,1550.5" in text.replace(" €", "").replace("1,550.50", "1550.5") or \
        "Total,\"1,550.50" in text


# --------------------------------------------------------------------------- #
# Showing numbers                                                             #
# --------------------------------------------------------------------------- #

GERMAN = NumberStyle(decimal=",", group=".")


@pytest.mark.parametrize("number, code, style, expected", [
    (1234.5, "General", NumberStyle(), "1234.5"),
    (1234.5, "General", GERMAN, "1234,5"),
    (0.1 + 0.2, "General", NumberStyle(), "0.3"),
    (1234.567, "0.00", NumberStyle(), "1234.57"),
    (1234.567, "#,##0.00", GERMAN, "1.234,57"),
    (-1234.5, "#,##0.00", NumberStyle(), "-1,234.50"),
    (0.256, "0%", NumberStyle(), "26%"),
    (0.256, "0.0 %", GERMAN, "25,6 %"),
    (1550.5, '#,##0.00 "€"', GERMAN, "1.550,50 €"),
    (1550.5, "#,##0.00 [$€-407]", GERMAN, "1.550,50 €"),
    (-5, '"$"#,##0.00;("$"#,##0.00)', NumberStyle(), "($5.00)"),
    (45306, "yyyy-mm-dd", NumberStyle(), "2024-01-15"),
    (45306, "dd.mm.yyyy", NumberStyle(), "15.01.2024"),
    (45306.5, "d mmm yyyy hh:mm", NumberStyle(), "15 Jan 2024 12:00"),
    (0.75, "h:mm AM/PM", NumberStyle(), "6:00 PM"),
    (12345678, "0.00E+00", NumberStyle(), "1.23E+07"),
    (3, "0", NumberStyle(), "3"),
])
def test_numbers_are_shown_in_their_format(number, code, style, expected):
    assert format_value(number, code, style) == expected


def test_values_of_other_kinds_are_shown_plainly():
    assert format_value(None, "0.00") == ""
    assert format_value("text", "0.00") == "text"
    assert format_value(True, "General") == "TRUE"
    assert format_value(sf.DIV0, "General") == "#DIV/0!"
    assert format_value(dt.date(2026, 10, 21), "General") == "2026-10-21"


def test_formats_are_told_apart():
    assert sheet_io.number_format_kind("0%") == "percent"
    assert sheet_io.number_format_kind('#,##0.00 "€"') == "currency"
    assert sheet_io.number_format_kind("dd/mm/yyyy") == "date"
    assert sheet_io.number_format_kind("hh:mm") == "time"
    assert sheet_io.number_format_kind("#,##0") == "number"
    assert not sheet_io.is_date_format('0.00 "m"')


# --------------------------------------------------------------------------- #
# Rows and columns on sheets full of formulas                                 #
# --------------------------------------------------------------------------- #

def test_only_formulas_whose_inputs_change_are_calculated_again(tmp_path):
    path = tmp_path / "counts.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in range(1, 6):
        ws.cell(row=row, column=1, value=row)
    ws["B1"] = "=COUNTBLANK(A1:A5)"
    ws["B2"] = "=B1*10"
    ws["B3"] = "=SUM({1,2})+A5"
    ws["B4"] = "=A1"
    wb.save(path)
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    assert value(sheet, "B1") == 0
    snapshot = book.insert(sheet, "row", 2, 2)
    assert sheet.cells[(0, 1)].formula == "=COUNTBLANK(A1:A7)"
    assert value(sheet, "B1") == 2 and value(sheet, "B2") == 20
    # A formula the engine cannot follow (an array constant) still moves.
    assert sheet.cells[(4, 1)].formula == "=SUM({1,2})+A7"
    assert sheet.cells[(5, 1)].formula == "=A1"
    book.restore(snapshot)
    assert value(sheet, "B1") == 0 and value(sheet, "B2") == 0
    book.delete(sheet, "row", 4, 1)
    assert sheet.cells[(0, 1)].formula == "=COUNTBLANK(A1:A4)"
    assert sheet.cells[(2, 1)].formula == "=SUM({1,2})+#REF!"
    edit(book, sheet, "A1", None)
    assert value(sheet, "B1") == 1 and value(sheet, "B2") == 10


def test_a_big_sheet_moves_every_formula_and_keeps_them_linked(tmp_path):
    path = tmp_path / "big.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in range(1, 601):
        ws.append([row, f"=A{row}*2", f"=SUM($B$1:B{row})"])
    wb.save(path)
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    total = value(sheet, "C600")
    book.insert(sheet, "row", 0, 1)
    assert sheet.cells[(600, 2)].formula == "=SUM($B$2:B601)"
    edit(book, sheet, "A2", 101)
    assert value(sheet, "C601") == total + 200


# --------------------------------------------------------------------------- #
# How cells look                                                              #
# --------------------------------------------------------------------------- #

def test_an_excel_cells_look_is_read_from_its_style(tmp_path):
    from tests.sheet_fixtures import styled_xlsx
    path = tmp_path / "Plan.xlsx"
    styled_xlsx(path)
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    header = book.look_at(sheet, 0, 0)
    assert header.fill == "#1F4E78" and header.color == "#FFFFFF" and header.bold
    assert header.align == "center"
    assert header.bottom == sheet_io.Border(2, "#000000")
    assert book.look_at(sheet, 4, 1).fill == "#FFFF00"
    red = book.look_at(sheet, 4, 2)
    assert red.color == "#C00000" and red.italic
    borders = book.look_at(sheet, 2, 2)
    assert borders.left == sheet_io.Border(1, "#808080") and borders.bottom.width == 3
    assert book.look_at(sheet, 2, 4).wrap
    # Accent 1 of the theme openpyxl writes (4F81BD), lightened by 60 %:
    # Excel's "Accent 1, Lighter 60%".
    assert book.look_at(sheet, 5, 0).fill == "#B9CDE5"
    assert book.look_at(sheet, 40, 8) == sheet_io.PLAIN
    assert sheet.frozen == (1, 1)
    assert sheet.heights[2] == pytest.approx(32 * 96 / 72)


def test_a_whole_row_style_colours_its_empty_cells(tmp_path):
    from openpyxl.styles import PatternFill
    path = tmp_path / "rows.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.row_dimensions[3].fill = PatternFill("solid", fgColor="00FF00")
    wb.save(path)
    book = load_workbook(str(path))
    assert book.look_at(book.sheets[0], 2, 7).fill == "#00FF00"
    book.insert(book.sheets[0], "row", 0, 1)
    assert book.look_at(book.sheets[0], 3, 7).fill == "#00FF00"


def test_an_ods_cells_look_follows_its_style_and_its_parents(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    budget = book.sheets[0]
    assert book.look_at(budget, 0, 0).bold
    # Yellow cells inside a run of repeated rows and cells, empty ones too.
    assert book.look_at(budget, 3, 0).fill == "#FFFF00"
    assert book.look_at(budget, 5, 1).fill == "#FFFF00"
    assert book.look_at(budget, 5, 2) == sheet_io.PLAIN
    assert book.look_at(budget, 900, 3) == sheet_io.PLAIN
    book.insert(budget, "row", 0, 1)
    assert book.look_at(budget, 4, 0).fill == "#FFFF00"
    assert book.look_at(budget, 7, 0) == sheet_io.PLAIN


def test_ods_borders_alignment_and_wrap_are_read():
    from lxml import etree
    styles_xml = etree.fromstring(
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"'
        ' xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0">'
        '<office:automatic-styles>'
        '<style:style style:name="base" style:family="table-cell">'
        '<style:table-cell-properties fo:border="0.06pt solid #000000"/>'
        '<style:text-properties fo:color="#0000ff"/></style:style>'
        '<style:style style:name="child" style:family="table-cell" '
        'style:parent-style-name="base">'
        '<style:table-cell-properties fo:border-left="2.5pt double #ff0000" '
        'fo:wrap-option="wrap" fo:background-color="transparent"/>'
        '<style:paragraph-properties fo:text-align="end"/></style:style>'
        '</office:automatic-styles></office:document-content>')
    looks = sheet_io._OdsLooks(sheet_io._OdsStyles(styles_xml, None))
    look = looks.look("child")
    assert look.left == sheet_io.Border(3, "#FF0000")
    assert look.top == sheet_io.Border(1, "#000000")
    assert look.align == "right" and look.wrap and look.fill is None
    assert look.color == "#0000FF"


# --------------------------------------------------------------------------- #
# OpenDocument: everything that names a cell follows it                       #
# --------------------------------------------------------------------------- #

_ODS_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
           'xmlns:calcext="urn:org:documentfoundation:names:experimental:calc:xmlns:calcext:1.0" '
           'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
           'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
           'xmlns:xlink="http://www.w3.org/1999/xlink"')


def _text_cell(text):
    return ('<table:table-cell office:value-type="string" calcext:value-type="string">'
            f'<text:p>{text}</text:p></table:table-cell>')


def _number_cell(number):
    return (f'<table:table-cell office:value-type="float" office:value="{number}" '
            f'calcext:value-type="float"><text:p>{number}</text:p></table:table-cell>')


def _formula_cell(formula, number):
    return (f'<table:table-cell table:formula="{formula}" office:value-type="float" '
            f'office:value="{number}" calcext:value-type="float"><text:p>{number}</text:p>'
            '</table:table-cell>')


def _ods(path, body, extra=None, settings=None):
    content = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {_ODS_NS} '
               'office:version="1.3"><office:body><office:spreadsheet>'
               f'{body}</office:spreadsheet></office:body></office:document-content>')
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(zipfile.ZipInfo("mimetype"), ODS_MIMETYPE, zipfile.ZIP_STORED)
        archive.writestr("content.xml", content, zipfile.ZIP_DEFLATED)
        if settings:
            archive.writestr("settings.xml", settings, zipfile.ZIP_DEFLATED)
        for name, data in (extra or {}).items():
            archive.writestr(name, data, zipfile.ZIP_DEFLATED)


def _content(path, name="content.xml"):
    with zipfile.ZipFile(path) as archive:
        return archive.read(name).decode("utf-8")


def test_a_row_change_on_a_sheet_with_a_space_in_its_name_moves_its_named_range(tmp_path):
    path = tmp_path / "costs.ods"
    rows = ("<table:table-row>" + _text_cell("Label") + _text_cell("Amount") + "</table:table-row>"
            "<table:table-row>" + _text_cell("a") + _number_cell(10) + "</table:table-row>"
            "<table:table-row>" + _text_cell("b") + _number_cell(20) + "</table:table-row>"
            "<table:table-row>" + _text_cell("x2") + _formula_cell("of:=[.B3]*2", 40) +
            "</table:table-row>")
    _ods(path, '<table:table table:name="Kosten 2026"><table:table-column '
               f'table:number-columns-repeated="2"/>{rows}</table:table>'
               '<table:named-expressions><table:named-range table:name="Amounts" '
               "table:base-cell-address=\"$'Kosten 2026'.$A$1\" "
               "table:cell-range-address=\"$'Kosten 2026'.$B$2:.$B$3\"/>"
               '</table:named-expressions>')
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    book.delete(sheet, "row", 1, 1)
    assert sheet.cells[(2, 1)].formula == "=B2*2" and value(sheet, "B3") == 40
    book.save()
    assert "table:cell-range-address=\"$'Kosten 2026'.$B$2:.$B$2\"" in _content(path)
    reread = load_workbook(str(path))
    assert reread.sheets[0].cells[(2, 1)].formula == "=B2*2"
    assert reread.backend.names["Amounts"] == "='Kosten 2026'!$B$2:$B$2"


def test_a_change_that_fails_halfway_leaves_everything_as_it_was(tmp_path, monkeypatch):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    book = load_workbook(str(path))
    budget = book.sheets[0]
    before_cells = dict(budget.cells)
    before_notes = dict(book.sheets[1].cells)
    calls = []
    real = sf.adjust_for_insert_delete

    def failing(*args, **kwargs):
        calls.append(1)
        if len(calls) == 3:
            raise RuntimeError("disk on fire")
        return real(*args, **kwargs)

    monkeypatch.setattr(sf, "adjust_for_insert_delete", failing)
    with pytest.raises(sheet_io.SheetFileError):
        book.delete(budget, "row", 1, 1)
    monkeypatch.setattr(sf, "adjust_for_insert_delete", real)
    assert budget.cells == before_cells
    assert book.sheets[1].cells == before_notes
    assert value(budget, "B7") == 1550.5
    book.save()
    reread = load_workbook(str(path))
    assert value(reread.sheets[0], "A2") == "Rent"
    assert reread.sheets[0].cells[(6, 1)].formula == "=SUM(B2:B5)"
    assert reread.backend.names["Amounts"] == "=Budget!$B$2:$B$5"


def test_ods_conditions_validation_and_named_expressions_follow_inserted_rows(tmp_path):
    path = tmp_path / "rules.ods"
    rows = "".join("<table:table-row>" + _text_cell("item") + _number_cell(r * 40) +
                   "<table:table-cell/>" + (_number_cell(100) if r == 1 else
                                            "<table:table-cell/>") + "</table:table-row>"
                   for r in range(1, 6))
    _ods(path,
         '<table:content-validations><table:content-validation table:name="val1" '
         'table:condition="of:cell-content-is-in-list([$S.$D$1:.$D$3])" '
         'table:base-cell-address="S.C2"/></table:content-validations>'
         f'<table:table table:name="S"><table:table-column table:number-columns-repeated="4"/>'
         f'{rows}<calcext:conditional-formats><calcext:conditional-format '
         'calcext:target-range-address="S.A2:S.A5"><calcext:condition '
         'calcext:apply-style-name="Bad" calcext:value="formula-is([.$B2]&gt;[.$D$1])" '
         'calcext:base-cell-address="S.A2"/></calcext:conditional-format>'
         '</calcext:conditional-formats></table:table>'
         '<table:named-expressions><table:named-expression table:name="Limit" '
         'table:base-cell-address="$S.$A$1" table:expression="of:=[$S.$D$1]*2"/>'
         '</table:named-expressions>')
    book = load_workbook(str(path))
    book.insert(book.sheets[0], "row", 0, 1)
    book.save()
    content = _content(path)
    assert 'calcext:value="formula-is([.$B3]&gt;[.$D$2])"' in content
    assert 'table:condition="of:cell-content-is-in-list([$S.$D$2:.$D$4])"' in content
    assert 'table:expression="of:=[$S.$D$2]*2"' in content
    assert 'calcext:base-cell-address="$S.A3"' in content


_CHART = (b'<?xml version="1.0" encoding="UTF-8"?><office:document-content '
          b'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          b'xmlns:chart="urn:oasis:names:tc:opendocument:xmlns:chart:1.0" '
          b'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"><office:body>'
          b'<office:chart><chart:chart chart:class="chart:bar"><chart:plot-area '
          b'table:cell-range-address="S.A1:S.B3"><chart:series '
          b'chart:values-cell-range-address="S.B1:S.B3" chart:label-cell-address="S.A1"/>'
          b'</chart:plot-area></chart:chart></office:chart></office:body>'
          b'</office:document-content>')

_SETTINGS = ('<?xml version="1.0" encoding="UTF-8"?><office:document-settings '
             'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
             'xmlns:config="urn:oasis:names:tc:opendocument:xmlns:config:1.0">'
             '<office:settings><config:config-item-set config:name="ooo:view-settings">'
             '<config:config-item-map-indexed config:name="Views"><config:config-item-map-entry>'
             '<config:config-item config:name="ActiveTable" config:type="string">S'
             '</config:config-item><config:config-item-map-named config:name="Tables">'
             '<config:config-item-map-entry config:name="S"><config:config-item '
             'config:name="VerticalSplitMode" config:type="short">2</config:config-item>'
             '<config:config-item config:name="VerticalSplitPosition" config:type="int">1'
             '</config:config-item></config:config-item-map-entry>'
             '</config:config-item-map-named></config:config-item-map-entry>'
             '</config:config-item-map-indexed></config:config-item-set></office:settings>'
             '</office:document-settings>')


def test_an_ods_charts_ranges_and_view_settings_follow_rename_and_rows(tmp_path):
    path = tmp_path / "chart.ods"
    rows = "".join("<table:table-row>" + _text_cell(f"k{r}") + _number_cell(r) +
                   "</table:table-row>" for r in range(1, 4))
    frame = ('<table:shapes><draw:frame svg:width="8cm" svg:height="5cm" svg:x="3cm" '
             'svg:y="1cm"><draw:object draw:notify-on-update-of-ranges="S.A1:S.B3" '
             'xlink:href="./Object 1" xlink:type="simple"/></draw:frame></table:shapes>')
    _ods(path, f'<table:table table:name="S">{frame}<table:table-column '
               f'table:number-columns-repeated="2"/>{rows}</table:table>',
         {"Object 1/content.xml": _CHART}, settings=_SETTINGS)
    book = load_workbook(str(path))
    book.rename_sheet(book.sheets[0], "Sales")
    book.insert(book.sheets[0], "row", 0, 1)
    book.save()
    chart = _content(path, "Object 1/content.xml")
    assert 'table:cell-range-address="$Sales.A2:.B4"' in chart
    assert 'chart:values-cell-range-address="$Sales.B2:.B4"' in chart
    assert 'chart:label-cell-address="$Sales.A2"' in chart
    assert 'draw:notify-on-update-of-ranges="$Sales.A2:.B4"' in _content(path)
    settings = _content(path, "settings.xml")
    assert 'config:config-item-map-entry config:name="Sales"' in settings
    assert ">Sales</config:config-item>" in settings
    # The frozen row moved down with the inserted one.
    assert load_workbook(str(path)).sheets[0].frozen == (2, 0)


def test_undo_of_a_rename_puts_charts_and_settings_back(tmp_path):
    path = tmp_path / "chart.ods"
    rows = "<table:table-row>" + _text_cell("k") + _number_cell(1) + "</table:table-row>"
    _ods(path, f'<table:table table:name="S"><table:table-column '
               f'table:number-columns-repeated="2"/>{rows}</table:table>',
         {"Object 1/content.xml": _CHART}, settings=_SETTINGS)
    book = load_workbook(str(path))
    book.restore(book.rename_sheet(book.sheets[0], "Sales"))
    book.save()
    assert 'table:cell-range-address="S.A1:S.B3"' in _content(path, "Object 1/content.xml")
    assert 'config:name="S"' in _content(path, "settings.xml")


# --------------------------------------------------------------------------- #
# Excel: formulas outside the cells, array formulas, results only Excel knows #
# --------------------------------------------------------------------------- #

def test_rules_validation_and_links_follow_inserted_rows_and_renamed_sheets(tmp_path):
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.hyperlink import Hyperlink
    path = tmp_path / "rules.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Entry"
    ws["H1"] = 100
    lists = wb.create_sheet("Lists")
    lists["A1"], lists["A2"], lists["B1"] = "yes", "no", 5
    ws.conditional_formatting.add("A2:A5", FormulaRule(formula=["$C2>$H$1"],
                                                       fill=PatternFill("solid",
                                                                        fgColor="FF0000")))
    ws.conditional_formatting.add("B2:B5", FormulaRule(formula=["B2>Lists!$B$1"],
                                                       fill=PatternFill("solid",
                                                                        fgColor="FF0000")))
    local = DataValidation(type="list", formula1="=$H$1:$H$3")
    local.add("C2:C5")
    other = DataValidation(type="list", formula1="=Lists!$A$1:$A$2")
    other.add("D2:D5")
    ws.add_data_validation(local)
    ws.add_data_validation(other)
    ws["E1"] = "go"
    ws["E1"].hyperlink = Hyperlink(ref="E1", location="Lists!A1", display="go")
    ws["F1"] = "here"
    ws["F1"].hyperlink = Hyperlink(ref="F1", location="Entry!H1", display="here")
    wb.save(path)
    book = load_workbook(str(path))
    book.insert(book.sheets[0], "row", 0, 1)
    book.rename_sheet(book.sheets[1], "Options")
    book.save()
    wb = openpyxl.load_workbook(path)
    entry = wb["Entry"]
    rules = {str(cf.sqref): [r.formula for r in cf.rules] for cf in entry.conditional_formatting}
    assert rules == {"A3:A6": [["$C3>$H$2"]], "B3:B6": [["B3>Options!$B$1"]]}
    assert sorted(dv.formula1 for dv in entry.data_validations.dataValidation) == \
        ["=$H$2:$H$4", "=Options!$A$1:$A$2"]
    assert entry["E2"].hyperlink.location == "Options!A1"
    assert entry["F2"].hyperlink.location == "Entry!H2"


def test_undo_puts_rules_validation_and_links_back(tmp_path):
    from openpyxl.formatting.rule import FormulaRule
    path = tmp_path / "rules.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.conditional_formatting.add("A1:A3", FormulaRule(formula=["$A1>Lists!$B$1"]))
    validation = DataValidation(type="list", formula1="=Lists!$A$1:$A$2")
    validation.add("B1:B3")
    ws.add_data_validation(validation)
    wb.create_sheet("Lists")
    wb.save(path)
    book = load_workbook(str(path))
    book.restore(book.rename_sheet(book.sheets[1], "Options"))
    book.restore(book.insert(book.sheets[0], "row", 0, 2))
    book.save()
    sheet = openpyxl.load_workbook(path).active
    assert [r.formula for cf in sheet.conditional_formatting for r in cf.rules] == \
        [["$A1>Lists!$B$1"]]
    assert sheet.data_validations.dataValidation[0].formula1 == "=Lists!$A$1:$A$2"


def test_an_array_formula_moves_as_an_array_formula(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    path = tmp_path / "arrays.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"], ws["B2"], ws["B3"] = 1, 2, 3
    ws["A5"] = ArrayFormula("A5:A7", "=B1:B3*2")
    ws["D1"] = ArrayFormula("D1", "=SUM(B1:B3*B1:B3)")
    wb.save(path)
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    book.insert(sheet, "row", 3, 1)       # between the inputs and the array
    book.insert(sheet, "row", 0, 1)       # above everything
    book.save()
    cells = openpyxl.load_workbook(path).active
    moved = cells["A7"].value
    assert isinstance(moved, ArrayFormula)
    assert moved.ref == "A7:A9" and moved.text == "=B2:B4*2"
    single = cells["D2"].value
    assert isinstance(single, ArrayFormula) and single.ref == "D2"
    assert single.text == "=SUM(B2:B4*B2:B4)"


def test_results_only_the_app_can_work_out_are_left_for_it(tmp_path):
    path = tmp_path / "stdev.ods"
    rows = ("<table:table-row>" + _number_cell(1) + _formula_cell("of:=STDEV([.A1:.A3])", 1) +
            _formula_cell("of:=[.B1]*100", 100) + _formula_cell("of:=[.A2]*2", 4) +
            "</table:table-row><table:table-row>" + _number_cell(2) +
            "</table:table-row><table:table-row>" + _number_cell(3) + "</table:table-row>")
    _ods(path, '<table:table table:name="S"><table:table-column '
               f'table:number-columns-repeated="4"/>{rows}</table:table>')
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    assert value(sheet, "B1") == 1 and sheet.cells[(0, 1)].state == sf.STATE_APP
    edit(book, sheet, "A3", 30)
    edit(book, sheet, "A2", 5)
    assert sheet.cells[(0, 1)].state == sf.STATE_STALE and value(sheet, "B1") is None
    assert sheet.cells[(0, 2)].state == sf.STATE_STALE
    assert value(sheet, "D1") == 10
    book.save()
    content = _content(path)
    # No result at all (LibreOffice calculates such a cell), never the old 1 or 100.
    assert '<table:table-cell table:formula="of:=STDEV([.A1:.A3])"/>' in content
    assert '<table:table-cell table:formula="of:=[.B1]*100"/>' in content
    assert 'office:value="10"' in content


def test_an_excel_result_only_excel_knows_is_not_written_stale(tmp_path):
    path = tmp_path / "unknown.xlsx"
    simple_xlsx(path)
    with zipfile.ZipFile(path) as archive:
        entries = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
        b'<c r="A3"><f>A1+A2</f><v></v></c>',
        b'<c r="A3"><f>_xlfn.FORECAST.ETS(A1,A1:A2,A1:A2)</f><v>7.5</v></c>')
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "A1", 50)
    book.save()
    assert openpyxl.load_workbook(path, data_only=True).active["A3"].value is None
    assert openpyxl.load_workbook(path).active["A3"].value == \
        "=_xlfn.FORECAST.ETS(A1,A1:A2,A1:A2)"


# --------------------------------------------------------------------------- #
# Excel: what a save keeps is known, everything else is named                 #
# --------------------------------------------------------------------------- #

def _rich_workbook(path):
    """A workbook with every feature the allow-list counts as kept."""
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.packaging.custom import StringProperty
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Amount"])
    for row in range(2, 8):
        ws.append([f"item {row}", row * 10])
    ws["D1"] = "=SUM(B2:B7)"
    ws["A2"].comment = Comment("Check", "Ann")
    ws["A3"].hyperlink = "https://example.org"
    ws.add_table(Table(displayName="Items", ref="A1:B7"))
    validation = DataValidation(type="whole", operator="greaterThan", formula1="0")
    validation.add("B2:B7")
    ws.add_data_validation(validation)
    ws.conditional_formatting.add("B2:B7", CellIsRule(operator="greaterThan", formula=["50"],
                                                      fill=PatternFill("solid",
                                                                       fgColor="00FF00")))
    ws.merge_cells("E1:F2")
    ws.freeze_panes = "A2"
    ws.print_options.gridLines = True
    ws.page_margins.left = 0.5
    ws.page_setup.orientation = "landscape"
    ws.oddHeader.center.text = "Budget"
    ws.row_breaks.append(Break(id=4))
    ws.col_breaks.append(Break(id=1))
    ws.protection.sheet = True
    ws.print_title_rows = "1:1"
    ws.sheet_properties.tabColor = "FF0000"
    other = wb.create_sheet("Other")
    other["A1"] = "=Data!B2*2"
    other.auto_filter.ref = "A1:A5"
    wb.defined_names["Amounts"] = DefinedName("Amounts", attr_text="Data!$B$2:$B$7")
    wb.security.lockStructure = True
    wb.custom_doc_props.append(StringProperty(name="Project", value="Budget"))
    wb.save(path)


def _top_elements(xml):
    root = sheet_io._sheet_head(xml)
    return {etree.QName(child).localname for child in root if isinstance(child.tag, str)}


def test_everything_the_allow_list_counts_as_kept_survives_a_save(tmp_path):
    path = tmp_path / "rich.xlsx"
    _rich_workbook(path)
    assert sheet_io.unsaveable_parts(str(path)) == []
    with zipfile.ZipFile(path) as archive:
        before = {name: archive.read(name) for name in archive.namelist()}
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "B2", 99)
    book.insert(book.sheets[0], "row", 7, 1)
    book.save()
    with zipfile.ZipFile(path) as archive:
        after = {name: archive.read(name) for name in archive.namelist()}
    derived = {"xl/calcChain.xml"}
    assert set(before) - derived <= set(after)
    for name in before:
        if re.fullmatch(r"xl/worksheets/sheet\d+\.xml|xl/workbook\.xml", name):
            assert _top_elements(before[name]) - {"dimension"} <= _top_elements(after[name]), name
    assert sheet_io.unsaveable_parts(str(path)) == []


def _add_to_package(path, changes):
    with zipfile.ZipFile(path) as archive:
        entries = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    for name, change in changes.items():
        entries[name] = change(entries.get(name, b""))
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)


_SLICERS = (b'<extLst><ext uri="{A8765BA9-456A-4dab-B4F3-ACF838C121DE}" '
            b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
            b'<x14:slicerList/></ext><ext uri="{01234567-89AB-4CDE-8F01-23456789ABCD}">'
            b"<unknown/></ext></extLst>")


# Printer settings, custom views, ignored error checks, sparklines, newer
# conditional formats and validation, threaded comments and add-ins are
# kept (tests/test_sheet_io.py: the producers' workbooks); these are not.
@pytest.mark.parametrize("changes, expected", [
    ({"xl/worksheets/sheet1.xml": lambda d: d.replace(b"</worksheet>",
                                                      _SLICERS + b"</worksheet>")},
     ["Slicers", "Newer Excel features ({01234567-89AB-4CDE-8F01-23456789ABCD})"]),
    ({"xl/metadata.xml": lambda d: b"<metadata/>"}, ["Dynamic array formulas"]),
    ({"xl/timelines/timeline1.xml": lambda d: b"<x/>"}, ["Timelines"]),
    ({"xl/ink/ink1.xml": lambda d: b"<x/>"}, ["Ink drawings"]),
    ({"xl/somethingNew/part1.xml": lambda d: b"<x/>"},
     ["Other parts of the workbook (xl/somethingNew/part1.xml)"]),
    ({"xl/worksheets/sheet1.xml": lambda d: d.replace(
        b"</worksheet>", b'<protectedRanges><protectedRange sqref="A1" name="r"/>'
                         b"</protectedRanges></worksheet>")}, ["Protected ranges"]),
    ({"xl/drawings/vmlDrawing9.vml": lambda d: b'<xml><x:ClientData ObjectType="Button"/></xml>'},
     ["Form controls"]),
])
def test_what_a_save_would_drop_is_named(tmp_path, changes, expected):
    path = tmp_path / "book.xlsx"
    simple_xlsx(path)
    _add_to_package(path, changes)
    assert sheet_io.unsaveable_parts(str(path)) == expected
    before = path.read_bytes()
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "C1", 1)
    with pytest.raises(sheet_io.UnsaveablePartsError):
        book.save()
    assert path.read_bytes() == before


def test_harmless_workbook_extensions_excel_writes_do_not_block_saving(tmp_path):
    path = tmp_path / "book.xlsx"
    simple_xlsx(path)
    _add_to_package(path, {"xl/workbook.xml": lambda d: d.replace(
        b"</workbook>", b'<extLst><ext uri="{140A7094-0E35-4892-8432-C4D2E57EDEB5}" '
                        b'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main">'
                        b'<x15:workbookPr chartTrackingRefBase="1"/></ext></extLst></workbook>')})
    assert sheet_io.unsaveable_parts(str(path)) == []


def test_a_save_by_another_app_while_the_file_is_read_is_noticed(tmp_path, monkeypatch):
    path = tmp_path / "race.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "version 1"
    wb.save(path)
    real = openpyxl.load_workbook

    def another_app_saves_meanwhile(*args, **kwargs):
        result = real(*args, **kwargs)
        if not kwargs.get("read_only"):
            newer = openpyxl.Workbook()
            newer.active["A1"] = "version 2"
            newer.save(path)
            os.utime(path, ns=(10**18, 10**18))
        return result

    monkeypatch.setattr(openpyxl, "load_workbook", another_app_saves_meanwhile)
    book = load_workbook(str(path))
    monkeypatch.setattr(openpyxl, "load_workbook", real)
    assert value(book.sheets[0], "A1") == "version 1"
    assert book.changed_on_disk()
    edit(book, book.sheets[0], "B1", "mine")
    with pytest.raises(sheet_io.ChangedOnDiskError):
        book.save()
    assert openpyxl.load_workbook(path).active["A1"].value == "version 2"


def test_a_file_only_touched_since_it_was_read_has_not_changed(styled_xlsx):
    book = load_workbook(styled_xlsx)
    os.utime(styled_xlsx, ns=(5, 5))
    assert not book.changed_on_disk()


TABLE_NS = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"


def _rows_of(path):
    root = etree.fromstring(_content(path).encode("utf-8"))
    table = root.find(f".//{TABLE_NS}table")
    rows = []
    for row in table.iter(f"{TABLE_NS}table-row"):
        rows.append([(cell.tag.split("}")[1], cell.get(f"{TABLE_NS}number-rows-spanned"),
                      "".join(cell.itertext())) for cell in row])
    return rows


def _merged_ods(path):
    rows = ('<table:table-row><table:table-cell table:number-rows-spanned="3" '
            'office:value-type="string"><text:p>Group</text:p></table:table-cell>' +
            _text_cell("r1") + "</table:table-row>"
            "<table:table-row><table:covered-table-cell/>" + _text_cell("r2") +
            "</table:table-row>"
            "<table:table-row><table:covered-table-cell/>" + _text_cell("r3") +
            "</table:table-row>"
            "<table:table-row>" + _text_cell("VISIBLE") + _text_cell("r4") +
            "</table:table-row>")
    _ods(path, '<table:table table:name="S"><table:table-column '
               f'table:number-columns-repeated="2"/>{rows}</table:table>')


def test_deleting_a_row_inside_an_ods_merge_shrinks_the_merge(tmp_path):
    path = tmp_path / "merge.ods"
    _merged_ods(path)
    book = load_workbook(str(path))
    book.delete(book.sheets[0], "row", 1, 1)
    book.save()
    rows = _rows_of(path)
    assert rows[0][0] == ("table-cell", "2", "Group")
    assert rows[1][0][0] == "covered-table-cell"
    assert rows[2][0] == ("table-cell", None, "VISIBLE")


def test_inserting_a_row_inside_an_ods_merge_grows_it_and_deleting_its_top_keeps_it(tmp_path):
    path = tmp_path / "merge.ods"
    _merged_ods(path)
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    book.insert(sheet, "row", 1, 1)
    assert sheet.merges == [(0, 0, 3, 0)]
    book.save()
    rows = _rows_of(path)
    assert rows[0][0] == ("table-cell", "4", "Group")
    assert [row[0][0] for row in rows[1:4]] == ["covered-table-cell"] * 3
    assert rows[4][0] == ("table-cell", None, "VISIBLE")
    book.delete(sheet, "row", 0, 1)
    book.save()
    rows = _rows_of(path)
    assert rows[0][0][:2] == ("table-cell", "3")
    assert rows[3][0] == ("table-cell", None, "VISIBLE")


def test_rows_inserted_next_to_groups_and_header_rows_stay_outside_them(tmp_path):
    path = tmp_path / "groups.ods"
    body = ('<table:table table:name="S"><table:table-column table:number-columns-repeated="2"/>'
            "<table:table-header-rows><table:table-row>" + _text_cell("Header") +
            "</table:table-row></table:table-header-rows>"
            '<table:table-row-group table:display="false">'
            '<table:table-row table:visibility="collapse">' + _text_cell("d1") +
            '</table:table-row><table:table-row table:visibility="collapse">' +
            _text_cell("d2") + "</table:table-row></table:table-row-group>"
            "<table:table-row>" + _text_cell("Total") + "</table:table-row></table:table>")
    _ods(path, body)
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    book.insert(sheet, "row", 3, 1)       # above Total, right after the group
    edit(book, sheet, "A4", "after group")
    book.insert(sheet, "row", 2, 1)       # between the two grouped rows
    edit(book, sheet, "A3", "inside group")
    book.insert(sheet, "row", 1, 1)       # right below the header rows
    edit(book, sheet, "A2", "after header")
    book.save()
    root = etree.fromstring(_content(path).encode("utf-8"))
    table = root.find(f".//{TABLE_NS}table")
    layout = []
    for child in table:
        tag = child.tag.split("}")[1]
        if tag == "table-row":
            layout.append(("row", child.get(f"{TABLE_NS}visibility"), "".join(child.itertext())))
        elif tag in ("table-header-rows", "table-row-group"):
            layout.append((tag, [(r.get(f"{TABLE_NS}visibility"), "".join(r.itertext()))
                                 for r in child]))
    assert layout == [
        ("table-header-rows", [(None, "Header")]),
        ("row", None, "after header"),
        ("table-row-group", [("collapse", "d1"), (None, "inside group"), ("collapse", "d2")]),
        ("row", None, "after group"),
        ("row", None, "Total"),
    ]


# --------------------------------------------------------------------------- #
# Saving as Excel, odd text, hidden sheets, frozen panes, tables              #
# --------------------------------------------------------------------------- #

def test_an_ods_saved_as_xlsx_keeps_its_names_and_its_date_system(tmp_path):
    path = tmp_path / "names.ods"
    date = ('<table:table-cell office:value-type="date" office:date-value="2026-10-21">'
            "<text:p>2026-10-21</text:p></table:table-cell>")
    later = ('<table:table-cell table:formula="of:=[.A2]+1" office:value-type="date" '
             'office:date-value="2026-10-22"><text:p>2026-10-22</text:p></table:table-cell>')
    rows = ("<table:table-row>" + _number_cell(10) + _formula_cell("of:=SUM(Amounts)*Rate", 6) +
            "</table:table-row><table:table-row>" + date + later + "</table:table-row>")
    _ods(path, '<table:calculation-settings><table:null-date table:date-value="1904-01-01"/>'
               '</table:calculation-settings><table:table table:name="S"><table:table-column '
               f'table:number-columns-repeated="2"/>{rows}</table:table>'
               '<table:named-expressions><table:named-range table:name="Amounts" '
               'table:base-cell-address="$S.$A$1" table:cell-range-address="$S.$A$1:.$A$1"/>'
               '<table:named-expression table:name="Rate" table:base-cell-address="$S.$A$1" '
               'table:expression="of:=0.6"/></table:named-expressions>')
    book = load_workbook(str(path))
    target = tmp_path / "names.xlsx"
    book.save(str(target))
    wb = openpyxl.load_workbook(target)
    assert wb.defined_names["Amounts"].attr_text == "S!$A$1:$A$1"
    assert wb.defined_names["Rate"].attr_text == "0.6"
    assert wb.epoch.year == 1904
    assert wb["S"]["A2"].value == dt.datetime(2026, 10, 21)
    cached = openpyxl.load_workbook(target, data_only=True)["S"]
    assert cached["B2"].value == dt.datetime(2026, 10, 22)
    assert cached["B1"].value == 6


def test_text_with_control_characters_is_saved_the_way_excel_writes_it(tmp_path):
    path = tmp_path / "ctrl.xlsx"
    simple_xlsx(path)
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "C1", "line1\x0bline2")
    book.save()
    with zipfile.ZipFile(path) as archive:
        assert b"line1_x000B_line2" in archive.read("xl/worksheets/sheet1.xml")


def test_the_last_sheet_that_is_not_hidden_cannot_be_deleted(tmp_path):
    path = tmp_path / "hidden.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Visible"
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    third = wb.create_sheet("Third")
    third["A1"] = 3
    wb.save(path)
    book = load_workbook(str(path))
    book.delete_sheet(book.sheets[0])
    with pytest.raises(sheet_io.SheetFileError):
        book.delete_sheet(book.sheets[1])
    book.save()
    saved = openpyxl.load_workbook(path)
    assert saved.sheetnames == ["Hidden", "Third"]
    assert saved.active.title == "Third"


def test_the_frozen_rows_in_the_file_match_what_the_sheet_shows(tmp_path):
    path = tmp_path / "frozen.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "head"
    wb.active.freeze_panes = "A2"
    wb.save(path)
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    book.insert(sheet, "row", 1, 1)
    assert sheet.frozen == (1, 0)
    book.insert(sheet, "row", 0, 1)
    assert sheet.frozen == (2, 0)
    book.save()
    assert openpyxl.load_workbook(path).active.freeze_panes == "A3"


def test_a_column_inserted_inside_an_excel_table_gets_a_header(tmp_path):
    from openpyxl.worksheet.table import Table
    path = tmp_path / "table.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item", "Amount"])
    ws.append(["a", 1])
    ws.add_table(Table(displayName="Items", ref="A1:B2"))
    wb.save(path)
    book = load_workbook(str(path))
    book.insert(book.sheets[0], "col", 1, 1)
    assert value(book.sheets[0], "B1") == "Column1"
    book.save()
    saved = openpyxl.load_workbook(path).active
    table = saved.tables["Items"]
    assert table.ref == "A1:C2"
    assert [c.name for c in table.tableColumns] == ["Item", "Column1", "Amount"]
    assert saved["B1"].value == "Column1"


# --------------------------------------------------------------------------- #
# Sheets read with a real XML parser                                          #
# --------------------------------------------------------------------------- #

def test_a_sheet_with_self_closing_cells_and_rows_opens(tmp_path):
    path = tmp_path / "styled.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["B1"].font = Font(bold=True)
    wb.active["A3"] = 3
    wb.save(path)
    _add_to_package(path, {"xl/worksheets/sheet1.xml": lambda d: re.sub(
        rb'<c r="B1" s="1"[^>]*>(?:</c>)?', b'<c r="B1" s="1"/>', d).replace(
        b'<row r="3"', b'<row r="2" ht="30" customHeight="1"/><row r="3"')})
    assert b'<c r="B1" s="1"/>' in zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml")
    assert sheet_io.unsaveable_parts(str(path)) == []
    book = load_workbook(str(path))
    assert value(book.sheets[0], "A3") == 3


def test_a_workbook_saved_here_opens_here_again(tmp_path):
    path = tmp_path / "merged.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active.merge_cells("B3:C4")
    wb.save(path)
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "A1", 2)
    book.save()
    again = load_workbook(str(path))
    assert value(again.sheets[0], "A1") == 2
    assert sheet_io.unsaveable_parts(str(path)) == []


def test_the_sheet_head_is_read_with_its_cells_left_out():
    xml = (b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
           b'<sheetData><row r="1"><c r="A1" s="1"/></row><row r="2"/></sheetData>'
           b'<pageMargins left="1"/><ignoredErrors><ignoredError sqref="A1"/></ignoredErrors>'
           b"</worksheet>")
    root = sheet_io._sheet_head(xml)
    assert [etree.QName(c).localname for c in root] == ["sheetData", "pageMargins",
                                                        "ignoredErrors"]
    assert len(root[0]) == 0


# --------------------------------------------------------------------------- #
# What a workbook can be saved as                                             #
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("kind, offered", [
    ("xlsx", [".xlsx", ".csv", ".tsv"]),
    ("xlsm", [".xlsm", ".csv", ".tsv"]),
    ("ods", [".ods", ".xlsx", ".csv", ".tsv"]),
    ("csv", [".csv", ".xlsx", ".tsv"]),
    ("tsv", [".tsv", ".xlsx", ".csv"]),
])
def test_save_as_offers_exactly_what_a_save_can_write(tmp_path, kind, offered):
    assert [f.ext for f in sheet_io.save_formats(kind)] == offered
    path = tmp_path / f"book.{kind}"
    if kind == "ods":
        realistic_ods(path)
    elif kind in ("csv", "tsv"):
        path.write_bytes(b"a,b\n1,2\n" if kind == "csv" else b"a\tb\n1\t2\n")
    else:
        simple_xlsx(path)
    book = load_workbook(str(path))
    for ext in (".xlsx", ".xlsm", ".ods", ".csv", ".tsv"):
        target = tmp_path / f"out{ext}"
        if ext in offered:
            book.save(str(target))
            assert target.exists()
        else:
            with pytest.raises(sheet_io.SheetFileError):
                book.save(str(target))
            assert not target.exists()
    book.close()


def test_the_file_name_follows_the_format_picked_in_the_dialog():
    assert sheet_io.save_as_path("/x/Budget.xlsx", ".csv", ".xlsx", "xlsx") == "/x/Budget.csv"
    assert sheet_io.save_as_path("/x/Budget.tsv", ".csv", ".xlsx", "xlsx") == "/x/Budget.tsv"
    assert sheet_io.save_as_path("/x/Budget", ".csv", ".xlsx", "xlsx") == "/x/Budget.csv"
    assert sheet_io.save_as_path("/x/Budget.ods", ".xlsx", ".xlsx", "xlsx") == \
        "/x/Budget.ods.xlsx"
    assert sheet_io.save_as_path("/x/Plan.ods", ".xlsx", ".ods", "ods") == "/x/Plan.xlsx"


def test_moving_a_sheet_keeps_the_sheet_the_file_opens_on(tmp_path):
    path = tmp_path / "three.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "One"
    wb.create_sheet("Two")
    wb.create_sheet("Three")
    wb.active = 1
    wb.save(path)
    book = load_workbook(str(path))
    book.move_sheet(book.sheets[2], 0)
    book.save()
    saved = openpyxl.load_workbook(path)
    assert saved.sheetnames == ["Three", "One", "Two"]
    assert saved.active.title == "Two"


def test_a_workbook_whose_file_was_deleted_is_not_saved_back_unasked(tmp_path):
    path = tmp_path / "gone.xlsx"
    simple_xlsx(path)
    book = load_workbook(str(path))
    edit(book, book.sheets[0], "A1", 9)
    os.remove(path)
    with pytest.raises(sheet_io.FileDeletedError):
        book.save()
    assert not path.exists()
    copy = tmp_path / "rescued.xlsx"
    book.save(str(copy))
    assert openpyxl.load_workbook(copy).active["A1"].value == 9


def test_typing_keeps_the_used_size_without_counting_every_cell(tmp_path):
    path = tmp_path / "size.xlsx"
    simple_xlsx(path)
    book = load_workbook(str(path))
    sheet = book.sheets[0]
    rows, cols = sheet.used_size()
    edit(book, sheet, "A1", 5)
    assert sheet._size == (rows, cols)            # known, not counted again
    edit(book, sheet, "J40", 1)
    assert sheet._size == (40, 10)
    edit(book, sheet, "J40", None)
    book.set_cells(sheet, {(39, 9): None})
    assert sheet.used_size() == (rows, cols)

