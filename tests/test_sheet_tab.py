# SPDX-FileCopyrightText: 2026 rinbal
# SPDX-License-Identifier: AGPL-3.0-or-later
"""The spreadsheet tab: keys as in Numbers and Excel, Undo, Paste, rows,
columns and sheets, what typing means, and saving that never loses data."""

import datetime as dt
import gc
import os
import threading
import time

import openpyxl
import pytest
import shiboken6
from PySide6.QtCore import QLocale, QMimeData, QPoint, Qt, QTimer
from PySide6.QtGui import QFontMetrics, QGuiApplication
from PySide6.QtTest import QTest
from PySide6.QtWidgets import (
    QApplication, QStyle, QStyleOptionHeader, QStyleOptionViewItem,
)

import sheet_formulas as sf
import sheet_io
import sheet_tab
from sheet_tab import SheetTab, parse_input
from tests.accessibility import unnamed_controls
from tests.sheet_fixtures import realistic_ods, simple_xlsx

ENGLISH = sheet_io.NumberStyle()
GERMAN = sheet_io.NumberStyle(decimal=",", group=".")
DE = QLocale(QLocale.Language.German, QLocale.Country.Germany)
US = QLocale(QLocale.Language.English, QLocale.Country.UnitedStates)


@pytest.fixture
def xlsx(tmp_path):
    path = tmp_path / "Budget.xlsx"
    simple_xlsx(path)
    return str(path)


@pytest.fixture
def tab(xlsx):
    tab = SheetTab(xlsx)
    tab._style, tab._locale = ENGLISH, US
    tab.resize(900, 600)
    tab.show()
    QApplication.processEvents()
    tab.grid.setFocus()
    yield tab
    tab.release()
    tab.close()
    tab.deleteLater()


def cell(tab, ref):
    return tab.sheet.cells.get(sf.parse_cell_name(ref))


def value(tab, ref):
    found = cell(tab, ref)
    return None if found is None else found.value


def type_into(tab, ref, text, finish=Qt.Key.Key_Return):
    """Type as a person does: the first key opens the cell's editor, which
    then has the keyboard."""
    tab.go_to(*sf.parse_cell_name(ref))
    QTest.keyClicks(tab.grid, text[0])
    assert tab._editor is not None
    QTest.keyClicks(tab._editor, text[1:])
    QTest.keyClick(tab._editor, finish)


# --------------------------------------------------------------------------- #
# Typing                                                                      #
# --------------------------------------------------------------------------- #

def test_the_tab_opens_on_the_first_sheet_at_a1(tab):
    assert tab.load_ok and tab.book is not None
    assert tab.current_cell() == (0, 0)
    assert tab.name_box.text() == "A1"
    assert tab.formula_field.text() == "1"
    assert tab.model.headerData(0, Qt.Orientation.Horizontal) == "A"
    assert tab.model.headerData(2, Qt.Orientation.Vertical) == "3"
    tab.go_to(2, 0)
    assert tab.formula_field.text() == "=A1+A2"
    assert tab.model.data(tab.model.index(2, 0)) == "3"


def test_typing_replaces_the_cell_and_return_moves_down(tab):
    type_into(tab, "A1", "41")
    assert value(tab, "A1") == 41
    assert tab.current_cell() == (1, 0)
    assert value(tab, "A3") == 43       # the formula followed
    assert tab.is_modified()


def test_tab_moves_right_and_shift_return_moves_up(tab):
    type_into(tab, "C2", "x", finish=Qt.Key.Key_Tab)
    assert value(tab, "C2") == "x" and tab.current_cell() == (1, 3)
    QTest.keyClick(tab.grid, Qt.Key.Key_Return, Qt.KeyboardModifier.ShiftModifier)
    assert tab.current_cell() == (0, 3)
    QTest.keyClick(tab.grid, Qt.Key.Key_Backtab, Qt.KeyboardModifier.ShiftModifier)
    assert tab.current_cell() == (0, 2)


def test_escape_cancels_an_edit(tab):
    tab.go_to(0, 0)
    QTest.keyClicks(tab.grid, "9")
    QTest.keyClicks(tab._editor, "99")
    QTest.keyClick(tab._editor, Qt.Key.Key_Escape)
    assert value(tab, "A1") == 1
    assert tab._editor is None
    assert not tab.is_modified()


def test_f2_edits_the_cell_in_place(tab):
    tab.go_to(2, 0)
    QTest.keyClick(tab.grid, Qt.Key.Key_F2)
    assert tab._editor is not None and tab._editor.text() == "=A1+A2"
    QTest.keyClicks(tab._editor, "*10")
    QTest.keyClick(tab._editor, Qt.Key.Key_Return)
    assert cell(tab, "A3").formula == "=A1+A2*10"
    assert value(tab, "A3") == 21


def test_the_formula_bar_commits_with_return_and_cancels_with_escape(tab):
    tab.go_to(3, 1)
    tab.formula_field.setFocus()
    tab.formula_field.setText("=sum(a1:a3)")
    QTest.keyClick(tab.formula_field, Qt.Key.Key_Return)
    assert cell(tab, "B4").formula == "=SUM(A1:A3)"
    assert value(tab, "B4") == 6
    tab.go_to(0, 0)
    tab.formula_field.setFocus()
    tab.formula_field.setText("123")
    tab.formula_field.setModified(True)
    QTest.keyClick(tab.formula_field, Qt.Key.Key_Escape)
    assert value(tab, "A1") == 1
    assert tab.formula_field.text() == "1"


def test_the_name_box_jumps_to_a_cell(tab):
    tab.name_box.setFocus()
    tab.name_box.setText("C12")
    QTest.keyClick(tab.name_box, Qt.Key.Key_Return)
    assert tab.current_cell() == (11, 2)
    tab.name_box.setText("AZ5000")
    QTest.keyClick(tab.name_box, Qt.Key.Key_Return)
    assert tab.current_cell() == (4999, 51)
    assert tab.model.rowCount() >= 5000


def test_delete_clears_the_selection_in_one_undoable_step(tab):
    tab.go_to(0, 0, end=(1, 1))
    QTest.keyClick(tab.grid, Qt.Key.Key_Delete)
    assert value(tab, "A1") is None and value(tab, "B1") is None
    assert value(tab, "A3") == 0
    tab.undo_stack.undo()
    assert value(tab, "A1") == 1 and value(tab, "B1") == "Name"
    assert value(tab, "A3") == 3


def test_undo_and_redo_walk_through_edits(tab):
    type_into(tab, "A1", "5")
    type_into(tab, "A2", "6")
    assert value(tab, "A3") == 11
    tab.undo_stack.undo()
    assert value(tab, "A2") == 2 and value(tab, "A3") == 7
    tab.undo_stack.undo()
    assert value(tab, "A1") == 1 and not tab.is_modified()
    tab.undo_stack.redo()
    tab.undo_stack.redo()
    assert value(tab, "A3") == 11


def test_a_circular_formula_shows_an_error_with_an_explanation(tab):
    type_into(tab, "A1", "=A3")
    assert value(tab, "A3") == sf.CIRCULAR
    index = tab.model.index(2, 0)
    assert tab.model.data(index) == "#CIRC!"
    assert tab.model.data(index, Qt.ItemDataRole.ToolTipRole)
    assert tab.model.data(index, Qt.ItemDataRole.ForegroundRole) is not None


def test_numbers_are_right_aligned_and_text_left(tab):
    right = int(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
    left = int(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
    assert tab.model.data(tab.model.index(0, 0), Qt.ItemDataRole.TextAlignmentRole) == right
    assert tab.model.data(tab.model.index(0, 1), Qt.ItemDataRole.TextAlignmentRole) == left


# --------------------------------------------------------------------------- #
# What typing means                                                           #
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("text, style, locale, expected_value, expected_format", [
    ("12,5", GERMAN, DE, 12.5, None),
    ("1.234,5", GERMAN, DE, 1234.5, None),
    ("12.5", ENGLISH, US, 12.5, None),
    ("1,234.5", ENGLISH, US, 1234.5, None),
    ("15 %", GERMAN, DE, 0.15, "0%"),
    ("2,5%", GERMAN, DE, 0.025, "0.0%"),
    ("1.234,50 €", GERMAN, DE, 1234.5, '#,##0.00 "€"'),
    ("$5.00", ENGLISH, US, 5, '"$"#,##0.00'),
    ("2026-10-21", ENGLISH, US, dt.date(2026, 10, 21), "yyyy-mm-dd"),
    ("21.10.2026", GERMAN, DE, dt.date(2026, 10, 21), "dd.mm.yyyy"),
    ("10/21/2026", ENGLISH, US, dt.date(2026, 10, 21), "m/d/yyyy"),
    ("14:30", ENGLISH, US, dt.time(14, 30), "hh:mm"),
    ("TRUE", ENGLISH, US, True, None),
    ("'0042", ENGLISH, US, "0042", None),
    ("hello", ENGLISH, US, "hello", None),
    ("007abc", ENGLISH, US, "007abc", None),
])
def test_what_is_typed_is_read_in_the_apps_language(text, style, locale, expected_value,
                                                    expected_format):
    parsed = parse_input(text, style=style, locale=locale)
    assert parsed.value == expected_value and type(parsed.value) is type(expected_value)
    assert parsed.number_format == expected_format


def test_a_typed_formula_is_written_in_capitals_outside_its_quotes():
    parsed = parse_input('=sum(a1:b2)&"abc"', style=ENGLISH, locale=US)
    assert parsed.formula == '=SUM(A1:B2)&"abc"'
    assert parse_input("=1+", style=ENGLISH, locale=US).error


def test_a_number_typed_into_a_percent_cell_is_a_percentage():
    parsed = parse_input("15", style=ENGLISH, current_format="0%", locale=US)
    assert parsed.value == 0.15


def test_text_that_looks_like_a_number_keeps_its_apostrophe_when_edited(tab):
    type_into(tab, "C1", "'0042")
    assert value(tab, "C1") == "0042"
    tab.go_to(0, 2)
    assert tab.formula_field.text() == "'0042"


def test_a_percent_typed_in_german_is_shown_and_saved_as_a_percent(tab, xlsx):
    tab._style, tab._locale = GERMAN, DE
    type_into(tab, "C1", "15 %")
    assert value(tab, "C1") == 0.15
    assert tab.model.data(tab.model.index(0, 2)) == "15%"
    assert tab.save().ok
    assert openpyxl.load_workbook(xlsx).active["C1"].number_format == "0%"


# --------------------------------------------------------------------------- #
# Clipboard                                                                   #
# --------------------------------------------------------------------------- #

def test_copy_puts_tab_separated_text_on_the_clipboard(tab):
    tab.go_to(0, 0, end=(2, 1))
    tab.copy()
    assert QGuiApplication.clipboard().text() == "1\tName\r\n2\t\r\n3\t\r\n"


def test_paste_of_a_block_fills_a_block(tab):
    data = QMimeData()
    data.setText("a\tb\r\nc\td\r\n")
    QGuiApplication.clipboard().setMimeData(data)
    tab.go_to(4, 0, end=(7, 3))
    tab.paste()
    assert [value(tab, r) for r in ("A5", "B5", "C5", "D5", "A6", "D8")] == \
        ["a", "b", "a", "b", "c", "d"]
    tab.undo_stack.undo()
    assert value(tab, "A5") is None


def test_pasted_numbers_are_read_like_typed_ones(tab):
    QGuiApplication.clipboard().setText("1.5\t20%\n")
    tab.go_to(5, 0)
    tab.paste()
    assert value(tab, "A6") == 1.5 and value(tab, "B6") == 0.2


def test_a_copied_formula_is_pasted_with_its_references_moved(tab):
    tab.go_to(2, 0)
    tab.copy()
    tab.go_to(2, 1)
    tab.paste()
    assert cell(tab, "B3").formula == "=B1+B2"
    tab.go_to(3, 0)
    tab.paste()
    assert cell(tab, "A4").formula == "=A2+A3"
    assert value(tab, "A4") == 5


def test_cut_empties_the_cells(tab):
    tab.go_to(0, 1)
    tab.cut()
    assert value(tab, "B1") is None
    assert "Name" in QGuiApplication.clipboard().text()


# --------------------------------------------------------------------------- #
# Rows, columns and sheets                                                    #
# --------------------------------------------------------------------------- #

def test_inserting_and_deleting_rows_adjusts_formulas(tab):
    assert tab.insert_rows(0, 1)
    assert cell(tab, "A4").formula == "=A2+A3" and value(tab, "A4") == 3
    assert tab.delete_rows(1, 1)
    assert cell(tab, "A3").formula == "=#REF!+A2"
    assert value(tab, "A3") == sf.REF
    tab.undo_stack.undo()
    tab.undo_stack.undo()
    assert cell(tab, "A3").formula == "=A1+A2" and value(tab, "A3") == 3


def test_inserting_and_deleting_columns(tab):
    tab.insert_columns(0, 2)
    assert value(tab, "C1") == 1 and value(tab, "D1") == "Name"
    assert cell(tab, "C3").formula == "=C1+C2"
    tab.delete_columns(0, 3)
    assert value(tab, "A1") == "Name"


def test_sheets_are_added_renamed_and_deleted(tab, monkeypatch):
    assert tab.add_sheet()
    assert tab.sheet_bar.count() == 2 and tab.current_sheet_index() == 1
    assert tab.sheet.name == "Sheet 2"
    assert tab.rename_sheet(1, "Plans")
    assert tab.sheet_bar.tabText(1) == "Plans"
    type_into(tab, "A1", "=Data!A3*2")
    assert value(tab, "A1") == 6
    assert tab.rename_sheet(0, "Numbers")
    assert tab.book.sheets[1].cells[(0, 0)].formula == "=Numbers!A3*2"
    import alerts
    monkeypatch.setattr(alerts, "confirm_destructive", lambda *a, **k: False)
    assert not tab.delete_sheet(0)
    monkeypatch.setattr(alerts, "confirm_destructive", lambda *a, **k: True)
    assert tab.delete_sheet(0)
    assert tab.sheet_bar.count() == 1
    assert tab.book.sheets[0].cells[(0, 0)].value == sf.REF
    tab.undo_stack.undo()
    assert tab.sheet_bar.count() == 2
    assert tab.book.sheets[1].cells[(0, 0)].value == 6


def test_a_sheet_is_renamed_in_place_by_double_clicking_its_tab(tab):
    tab.start_rename(0)
    assert tab._rename_edit is not None
    tab._rename_edit.setText("Ledger")
    QTest.keyClick(tab._rename_edit, Qt.Key.Key_Return)
    assert tab.sheet_bar.tabText(0) == "Ledger"


def test_a_column_width_is_undoable_and_saved(tab, xlsx):
    tab.grid.horizontalHeader().resizeSection(1, 180)
    digit = tab._digit_width()
    assert tab.sheet.widths[1] == pytest.approx(sheet_tab._file_width(180, digit))
    assert tab.save().ok
    # Excel's unit: characters of the widest digit, with its padding.
    assert openpyxl.load_workbook(xlsx).active.column_dimensions["B"].width == \
        pytest.approx((180 - sheet_tab._padding(digit)) / digit, abs=0.01)
    tab.undo_stack.undo()
    assert 1 not in tab.sheet.widths


# --------------------------------------------------------------------------- #
# Saving                                                                      #
# --------------------------------------------------------------------------- #

def test_save_writes_the_file_and_clears_the_modified_state(tab, xlsx):
    seen = []
    tab.modification_changed.connect(seen.append)
    type_into(tab, "A1", "10")
    assert seen == [True]
    result = tab.save()
    assert result.ok and result.status == sheet_tab.SAVED
    assert seen == [True, False]
    assert openpyxl.load_workbook(xlsx, data_only=True).active["A3"].value == 12


def test_a_file_open_in_another_app_is_read_only_until_it_closes(tab, xlsx):
    lock = os.path.join(os.path.dirname(xlsx), ".~lock.Budget.xlsx#")
    open(lock, "w", encoding="utf-8").close()
    tab._check_lock()
    assert tab.read_only and tab._notice.isVisible()
    assert not tab.commit_text(0, 0, "5")
    os.remove(lock)
    tab._check_lock()
    assert not tab.read_only
    type_into(tab, "A1", "5")
    open(lock, "w", encoding="utf-8").close()
    result = tab.save()
    assert result.status == sheet_tab.IN_USE
    os.remove(lock)


def test_saving_a_file_that_changed_on_disk_asks_first(tab, xlsx):
    type_into(tab, "A1", "5")
    wb = openpyxl.load_workbook(xlsx)
    wb.active["C9"] = "saved by another app"
    wb.save(xlsx)
    result = tab.save()
    assert result.status == sheet_tab.CHANGED_ON_DISK
    assert tab.save(overwrite_changes=True).ok


def test_a_file_only_touched_on_disk_has_not_changed(tab, xlsx):
    type_into(tab, "A1", "5")
    os.utime(xlsx, ns=(1, 1))
    assert tab.save().ok


def test_a_workbook_with_charts_asks_for_a_copy(tmp_path):
    import zipfile
    plain = tmp_path / "plain.xlsx"
    simple_xlsx(plain)
    path = tmp_path / "Chart.xlsx"
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(path, "w") as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr("xl/charts/chart1.xml", b"<c/>")
    before = path.read_bytes()
    tab = SheetTab(str(path))
    try:
        assert tab._notice.text()
        tab.commit_text(0, 0, "7")
        result = tab.save()
        assert result.status == sheet_tab.NEEDS_COPY and result.parts == ("Charts",)
        assert path.read_bytes() == before
        copy = tmp_path / "Chart copy.xlsx"
        result = tab.save_as(str(copy))
        assert result.ok and result.switched
        assert tab._file_path == str(copy) and not tab.is_modified()
        assert openpyxl.load_workbook(copy).active["A1"].value == 7
        assert tab.save().ok
    finally:
        tab.release()


def test_an_unmodified_tab_reloads_silently_and_a_modified_one_asks(tab, xlsx):
    wb = openpyxl.load_workbook(xlsx)
    wb.active["A1"] = 100
    wb.save(xlsx)
    os.utime(xlsx, ns=(5, 5))
    tab._do_scheduled_reload()
    assert value(tab, "A1") == 100
    type_into(tab, "B2", "mine")
    wb.active["A1"] = 200
    wb.save(xlsx)
    os.utime(xlsx, ns=(9, 9))
    tab._do_scheduled_reload()
    assert not tab._changed_bar.isHidden()
    assert value(tab, "A1") == 100
    tab._on_keep_clicked()
    assert tab.save().ok
    assert openpyxl.load_workbook(xlsx).active["B2"].value == "mine"


def test_save_as_csv_exports_the_sheet_and_keeps_the_workbook(tab, tmp_path):
    target = tmp_path / "export.csv"
    result = tab.save_as(str(target))
    assert result.ok and not result.switched
    assert target.read_bytes().decode("utf-8-sig").startswith("1,Name\r\n2,")
    assert tab._file_path.endswith("Budget.xlsx")


def test_an_ods_is_edited_and_saved_in_place(tmp_path):
    path = tmp_path / "Budget.ods"
    realistic_ods(path)
    tab = SheetTab(str(path))
    tab._style, tab._locale = ENGLISH, US
    try:
        assert tab.sheet_bar.count() == 2
        assert tab.model.data(tab.model.index(1, 1)) == "1.200,00 €"
        tab.commit_text(2, 1, "400")
        assert tab.save().ok
        reread = sheet_io.load_workbook(str(path))
        assert reread.sheets[0].cells[(6, 1)].value == 1600
    finally:
        tab.release()


def test_a_csv_keeps_what_is_typed_as_text(tmp_path):
    path = tmp_path / "list.csv"
    path.write_bytes(b"a;b\r\n1;2\r\n")
    tab = SheetTab(str(path))
    try:
        assert not tab._add_button.isVisibleTo(tab)
        tab.commit_text(1, 1, "=1+1")
        tab.commit_text(2, 0, "12,5")
        assert tab.sheet.cells[(1, 1)].value == "=1+1"
        assert tab.save().ok
        assert path.read_bytes() == b"a;b\r\n1;=1+1\r\n12,5;\r\n"
    finally:
        tab.release()


def test_a_large_file_loads_in_the_background(tmp_path):
    path = tmp_path / "big.csv"
    path.write_text("".join(f"{i},row {i}\n" for i in range(3000)), encoding="utf-8")
    tab = SheetTab(str(path), background=True)
    try:
        assert tab.is_loading()
        assert tab._stack.currentWidget() is tab._loading
        for _ in range(500):
            QApplication.processEvents()
            if not tab.is_loading():
                break
            QTest.qWait(10)
        assert tab.book is not None
        assert tab.sheet.cells[(2999, 1)].value == "row 2999"
    finally:
        tab.release()


def test_a_damaged_file_shows_why(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"nothing")
    tab = SheetTab(str(path))
    try:
        assert not tab.load_ok and tab.load_error
        assert tab._stack.currentWidget() is tab._failure
        assert not tab.save().ok
    finally:
        tab.release()


# --------------------------------------------------------------------------- #
# The bar on top and accessibility                                            #
# --------------------------------------------------------------------------- #

def test_the_open_button_names_the_app(tab):
    opened, asked = [], []
    tab.open_requested.connect(lambda: opened.append(True))
    tab.open_with_requested.connect(asked.append)
    tab.set_default_app("Numbers")
    assert tab._bar.open_button.text() == "Open in Numbers"
    tab._bar.open_button.click()
    tab._bar.open_menu_button.click()
    assert opened == [True] and isinstance(asked[0], QPoint)
    tab.set_default_app(None)
    assert tab._bar.open_button.text() == "Open With…"
    tab._bar.open_button.click()
    assert len(asked) == 2 and opened == [True]


def test_every_control_has_a_name_for_screen_readers(tab):
    tab.start_rename(0)
    tab.go_to(0, 0)
    QTest.keyClick(tab.grid, Qt.Key.Key_F2)
    assert unnamed_controls(tab) == []
    assert tab.grid.accessibleName() and tab.name_box.accessibleName()
    assert tab.formula_field.accessibleName() and tab.sheet_bar.accessibleName()


def test_the_title_is_the_file_name(tab):
    assert tab.document_title() == "Budget"
    assert tab._info_label.text().startswith("Excel Spreadsheet · ")


def test_a_german_tab_shows_and_reads_formulas_in_german(tab, xlsx):
    tab._style, tab._locale, tab._language = GERMAN, DE, "de"
    type_into(tab, "B1", "=RUNDEN(A1*1,19;2)")
    assert cell(tab, "B1").formula == "=ROUND(A1*1.19,2)"
    assert value(tab, "B1") == 1.19
    tab.go_to(0, 1)
    assert tab.formula_field.text() == "=RUNDEN(A1*1,19;2)"
    tab.go_to(2, 0)
    assert tab.formula_field.text() == "=A1+A2"
    type_into(tab, "C1", "=SUM(A1:A2;10)")
    assert value(tab, "C1") == 13
    type_into(tab, "D1", "=KONKATENIEREN(\"a;b\";WAHR)")
    assert cell(tab, "D1").formula == '=CONCAT("a;b",TRUE)'
    assert tab.save().ok
    ws = openpyxl.load_workbook(xlsx).active
    assert ws["B1"].value == "=ROUND(A1*1.19,2)"
    assert ws["D1"].value == '=_xlfn.CONCAT("a;b",TRUE)'


def test_cells_copied_in_german_paste_with_the_same_formula(tab):
    tab._style, tab._locale, tab._language = GERMAN, DE, "de"
    type_into(tab, "B1", "=SUMME(A1;1,5)")
    tab.go_to(0, 1)
    tab.copy()
    tab.go_to(1, 1)
    tab.paste()
    assert cell(tab, "B2").formula == "=SUM(A2,1.5)"


# --------------------------------------------------------------------------- #
# Looks and frozen panes                                                      #
# --------------------------------------------------------------------------- #

@pytest.fixture
def styled(tmp_path):
    from tests.sheet_fixtures import styled_xlsx
    path = tmp_path / "Plan.xlsx"
    styled_xlsx(path)
    tab = SheetTab(str(path))
    tab._style, tab._locale = ENGLISH, US
    tab.resize(800, 500)
    tab.show()
    QApplication.processEvents()
    tab.grid.setFocus()
    yield tab
    tab.release()
    tab.close()


def test_the_grid_draws_fills_text_colours_fonts_and_alignment(styled):
    model = styled.model
    assert model.data(model.index(4, 1), Qt.ItemDataRole.BackgroundRole).name() == "#ffff00"
    assert model.data(model.index(4, 2), Qt.ItemDataRole.ForegroundRole).name() == "#c00000"
    assert model.data(model.index(4, 2), Qt.ItemDataRole.FontRole).italic()
    header = model.index(0, 1)
    assert model.data(header, Qt.ItemDataRole.ForegroundRole).name() == "#ffffff"
    assert model.data(header, Qt.ItemDataRole.FontRole).bold()
    assert model.data(header, Qt.ItemDataRole.TextAlignmentRole) == \
        int(Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter)
    assert styled.grid.rowHeight(2) == round(32 * 96 / 72)
    # Painting with borders and wrapped text works on screen.
    styled.grid.viewport().repaint()


def test_unreadable_text_colours_give_way_on_the_dark_theme(styled):
    styled.update_theme(True)
    look = sheet_io.Look(color="#111111")
    assert styled.text_color(look) is None
    assert styled.text_color(sheet_io.Look(fill="#FFFF00")).name() == "#000000"
    border = sheet_io.Border(1, None)
    assert styled.border_color(border, sheet_io.PLAIN).name() != "#000000"


def test_frozen_rows_and_columns_stay_while_the_grid_scrolls(styled):
    panes = styled._panes
    assert styled._frozen == (1, 1)
    assert all(not panes[kind].isHidden() for kind in ("rows", "cols", "corner"))
    styled.go_to(45, 9)
    QApplication.processEvents()
    grid = styled.grid
    assert grid.verticalScrollBar().value() > 0
    assert panes["rows"].verticalScrollBar().value() == 0
    assert panes["cols"].verticalScrollBar().value() == grid.verticalScrollBar().value()
    assert panes["rows"].horizontalScrollBar().value() == grid.horizontalScrollBar().value()
    assert panes["corner"].verticalScrollBar().value() == 0
    # The pane covers the header and exactly the frozen row.
    assert panes["rows"].height() == grid.horizontalHeader().height() + grid.rowHeight(0)
    # The current cell is never hidden under the frozen row.
    rect = grid.visualRect(styled.model.index(45, 9))
    assert rect.top() >= grid.rowHeight(0)


def test_a_frozen_cell_is_edited_in_its_pane(styled):
    styled.go_to(0, 2)
    QTest.keyClicks(styled.grid, "Z")
    assert styled._editor is not None
    assert styled._editor_view is styled._panes["rows"]
    QTest.keyClicks(styled._editor, "1")
    QTest.keyClick(styled._editor, Qt.Key.Key_Return)
    assert value(styled, "C1") == "Z1"
    assert styled.current_cell() == (1, 2)


def test_widening_a_column_on_a_frozen_pane_widens_the_grid(styled):
    styled._panes["rows"].horizontalHeader().resizeSection(2, 150)
    assert styled.grid.columnWidth(2) == 150
    assert styled.sheet.widths[2] == pytest.approx(
        sheet_tab._file_width(150, styled._digit_width()))
    styled.undo_stack.undo()
    assert styled._panes["rows"].columnWidth(2) == styled.grid.columnWidth(2)


def test_frozen_panes_are_named_for_screen_readers(styled):
    assert unnamed_controls(styled) == []


def test_a_row_change_that_fails_says_so_and_leaves_no_undo_step(tab, monkeypatch):
    import alerts
    told = []
    monkeypatch.setattr(alerts, "inform", lambda *a, **k: told.append(k.get("title")))

    def broken(*args, **kwargs):
        raise RuntimeError("broken backend")

    before = dict(tab.sheet.cells)
    monkeypatch.setattr(tab.book.backend, "shift", broken)
    assert not tab.insert_rows(0, 1)
    assert told and tab.undo_stack.count() == 0
    assert tab.sheet.cells == before
    assert not tab.is_modified()


def test_control_characters_typed_or_pasted_are_dropped(tab):
    assert tab.commit_text(0, 2, "a\x0bb")
    assert value(tab, "C1") == "ab"
    QGuiApplication.clipboard().setText("x\x01y\tz\n")
    tab.go_to(4, 0)
    tab.paste()
    assert value(tab, "A5") == "xy" and value(tab, "B5") == "z"
    assert tab.save().ok


def test_text_longer_than_a_cell_holds_is_refused(tab, monkeypatch):
    import alerts
    told = []
    monkeypatch.setattr(alerts, "inform", lambda *a, **k: told.append(k.get("title")))
    assert not tab.commit_text(5, 0, "x" * (sheet_tab.MAX_CELL_TEXT + 1))
    assert told and value(tab, "A6") is None


def test_formulas_that_spill_over_several_cells_are_refused(tab, monkeypatch):
    import alerts
    told = []
    monkeypatch.setattr(alerts, "inform", lambda *a, **k: told.append(k.get("title")))
    assert not tab.commit_text(5, 0, "=SORT(A1:A3)")
    assert told and "SORT" in told[0]
    tab._language = "de"
    assert not tab.commit_text(5, 0, "=FILTER(A1:A3;A1:A3>1)")
    assert value(tab, "A6") is None


def test_a_save_that_fails_unexpectedly_reports_instead_of_raising(tab, monkeypatch):
    type_into(tab, "A1", "5")

    def broken(*args, **kwargs):
        raise ValueError("something nobody expected")

    monkeypatch.setattr(tab.book.backend, "serialize", broken)
    result = tab.save()
    assert result.status == sheet_tab.FAILED and result.message
    assert tab.is_modified()


# --------------------------------------------------------------------------- #
# Loading in the background                                                   #
# --------------------------------------------------------------------------- #

def _slow_loading(monkeypatch, *, read_first):
    """load_workbook that waits for ``gate`` before (or after) reading."""
    import threading
    gate = threading.Event()
    real = sheet_io.load_workbook

    def slow(path, progress=None):
        if read_first:
            book = real(path, progress)
            gate.wait(10)
            return book
        gate.wait(10)
        return real(path, progress)

    monkeypatch.setattr(sheet_io, "load_workbook", slow)
    return gate


def _wait_loaded(tab):
    for _ in range(1000):
        QApplication.processEvents()
        if not tab.is_loading():
            return
        QTest.qWait(5)


@pytest.mark.parametrize("read_first", [True, False])
def test_a_sheet_renamed_while_it_loads_saves_under_its_new_name(tmp_path, monkeypatch,
                                                                 read_first):
    old = tmp_path / "Budget.xlsx"
    new = tmp_path / "Budget 2026.xlsx"
    simple_xlsx(old)
    gate = _slow_loading(monkeypatch, read_first=read_first)
    tab = SheetTab(str(old), background=True)
    try:
        QTest.qWait(50)
        os.rename(old, new)
        tab.follow_move(str(new))
        gate.set()
        _wait_loaded(tab)
        assert tab.load_ok and tab.book is not None
        assert tab.book.path == str(new)
        assert tab.commit_text(0, 0, "7")
        assert tab.save().ok
        assert not old.exists()
        assert openpyxl.load_workbook(new).active["A1"].value == 7
    finally:
        gate.set()
        tab.release()


def test_a_tab_closed_while_its_sheet_loads_lets_the_worker_end_quietly(tmp_path,
                                                                        monkeypatch):
    import threading
    path = tmp_path / "Budget.xlsx"
    simple_xlsx(path)
    gate = _slow_loading(monkeypatch, read_first=True)
    problems = []
    monkeypatch.setattr(threading, "excepthook", lambda args: problems.append(args))
    tab = SheetTab(str(path), background=True)
    worker = next(t for t in threading.enumerate() if t.name == "sheet-loader")
    tab.release()
    import shiboken6
    shiboken6.delete(tab)                 # the tab, and its loader, are gone
    gate.set()
    worker.join(10)
    assert not worker.is_alive()
    assert problems == []


def test_a_file_moved_to_the_trash_is_shown_as_gone_and_never_saved_back(tab, xlsx):
    type_into(tab, "A1", "5")
    os.remove(xlsx)
    tab._do_scheduled_reload()
    assert tab.is_deleted()
    assert tab._changed_label.text() == "This file was moved to the Trash or deleted."
    result = tab.save()
    assert result.status == sheet_tab.DELETED
    assert not os.path.exists(xlsx)
    closing = []
    tab.close_requested.connect(lambda: closing.append(True))
    tab._close_button.click()
    assert closing == [True]


def test_a_file_put_back_from_the_trash_is_its_file_again(tab, xlsx, tmp_path):
    kept = tmp_path / "kept.xlsx"
    os.rename(xlsx, kept)
    tab._do_scheduled_reload()
    assert tab.is_deleted()
    os.rename(kept, xlsx)
    tab._do_scheduled_reload()
    assert not tab.is_deleted() and tab._changed_bar.isHidden()
    type_into(tab, "A1", "6")
    assert tab.save().ok



def test_the_sheet_tab_has_the_bar_every_file_tab_has(tab, xlsx):
    from viewer_tabs import FileBar, lock_notice_text
    bar = tab._bar
    assert isinstance(bar, FileBar)
    assert bar.kind_icon.pixmap() is not None and not bar.kind_icon.pixmap().isNull()
    assert bar.name_label.full_text() == "Budget.xlsx"
    assert bar.facts_label.full_text().startswith("1 sheet \u00b7 ")
    tab.set_default_app("Numbers")
    for button in (bar.open_menu_button, bar.reveal_button):
        assert button.text() == "" and not button.icon().isNull()
    tab.set_reveal_label("Show in Explorer")
    assert bar.reveal_button.accessibleName() == "Show in Explorer"
    lock = os.path.join(os.path.dirname(xlsx), "~$Budget.xlsx")
    open(lock, "w", encoding="utf-8").close()
    tab._check_lock()
    assert tab._notice.text() == lock_notice_text() == (
        "This file is open in another app. You are seeing its last saved version.")
    os.remove(lock)


def test_select_all_on_a_long_sheet_never_walks_its_cells(xlsx, monkeypatch):
    # Qt's headers asked about every cell of each column shown: after
    # Select All on 50 000 rows a repaint walked 400 000 cells.
    asked = []
    flags = sheet_tab.SheetModel.flags
    monkeypatch.setattr(sheet_tab.SheetModel, "flags",
                        lambda self, index: asked.append(1) or flags(self, index))
    tab = SheetTab(xlsx)
    try:
        tab.resize(900, 600)
        tab.show()
        tab.model.ensure(49_999, 5)
        QApplication.processEvents()
        asked.clear()
        tab.grid.selectAll()
        tab.grid.horizontalHeader().viewport().repaint()
        tab.grid.verticalHeader().viewport().repaint()
        QApplication.processEvents()
        assert len(asked) < 5_000
        header = tab.grid.horizontalHeader()
        option = QStyleOptionHeader()
        header.initStyleOptionForIndex(option, 2)
        assert option.state & QStyle.StateFlag.State_Sunken
        tab.go_to(3, 1, end=(5, 2))
        option = QStyleOptionHeader()
        header.initStyleOptionForIndex(option, 2)
        assert option.state & QStyle.StateFlag.State_On
        assert not option.state & QStyle.StateFlag.State_Sunken
        option = QStyleOptionHeader()
        header.initStyleOptionForIndex(option, 4)
        assert not option.state & QStyle.StateFlag.State_On
    finally:
        # Gone before the patched flags are: Qt keeps calling what it found.
        tab.release()
        shiboken6.delete(tab)



# --------------------------------------------------------------------------- #
# Large changes on a worker thread                                            #
# --------------------------------------------------------------------------- #

@pytest.fixture
def every_change_is_large(monkeypatch):
    """Every change runs as one on a few hundred thousand cells does: on a
    worker thread while the tab waits."""
    monkeypatch.setattr(sheet_tab, "BACKGROUND_CELLS", -1)


def test_a_large_change_runs_on_a_worker_thread_and_undoes(tab, every_change_is_large,
                                                           monkeypatch):
    threads = []
    insert = tab.book.insert
    monkeypatch.setattr(tab.book, "insert", lambda *a, **k: threads.append(
        threading.current_thread()) or insert(*a, **k))
    type_into(tab, "A1", "5")
    assert value(tab, "A1") == 5
    assert tab.insert_rows(0)
    assert threads and threads[0] is not threading.main_thread()
    assert value(tab, "A2") == 5 and value(tab, "A1") is None
    tab.undo_stack.undo()
    assert value(tab, "A1") == 5
    tab.undo_stack.undo()
    assert value(tab, "A1") != 5
    assert not tab.is_busy() and not tab.read_only
    assert tab._stack.currentWidget() is tab.grid and tab._job_progress.isHidden()


def test_while_a_large_change_runs_the_tab_waits_and_shows_the_sheet(tab, every_change_is_large,
                                                                    monkeypatch):
    type_into(tab, "A1", "7")
    release = threading.Event()
    save = tab.book.save
    monkeypatch.setattr(tab.book, "save", lambda *a, **k: release.wait(5) and save(*a, **k))
    seen = {}

    def look():
        seen.update(busy=tab.is_busy(), read_only=tab.read_only,
                    picture=tab._stack.currentWidget() is tab._picture,
                    can_undo=tab.undo_stack.canUndo(), typed=tab.commit_text(1, 1, "x"),
                    info=tab._info_label.text(), progress=tab._job_progress.isVisible())
        tab.undo_stack.undo()           # waits for the save: nothing happens
        seen["a1"] = value(tab, "A1")
        release.set()

    QTimer.singleShot(50, look)
    assert tab.save().ok
    assert seen == dict(busy=True, read_only=True, picture=True, can_undo=False, typed=False,
                        info="Saving…", progress=True, a1=7)
    assert not tab.is_busy() and not tab.read_only and tab.undo_stack.canUndo()
    assert tab._stack.currentWidget() is tab.grid
    assert openpyxl.load_workbook(tab._file_path).active["A1"].value == 7


def test_a_save_asked_for_during_a_save_waits_for_it(tab, every_change_is_large, monkeypatch):
    type_into(tab, "A1", "8")
    order, second = [], []
    release = threading.Event()
    save = tab.book.save

    def slow_save(*a, **k):
        order.append("start")
        if len(order) == 1:
            release.wait(5)
        result = save(*a, **k)
        order.append("end")
        return result

    monkeypatch.setattr(tab.book, "save", slow_save)
    QTimer.singleShot(30, lambda: second.append(tab.save()))
    QTimer.singleShot(150, release.set)
    assert tab.save().ok
    assert second and second[0].ok
    assert order == ["start", "end", "start", "end"]


def test_a_large_change_that_cannot_be_made_says_why(tab, every_change_is_large, monkeypatch):
    told = []
    monkeypatch.setattr("alerts.inform", lambda *a, **k: told.append(k.get("title")))

    def refuse(*_a, **_k):
        raise sheet_io.SheetFileError("Not here.")

    monkeypatch.setattr(tab.book, "insert", refuse)
    assert not tab.insert_rows(0)
    assert told == ["Not here."]
    assert not tab.undo_stack.canUndo() and not tab.is_busy()


def test_large_copies_clears_and_pastes_match_small_ones(tab, monkeypatch):
    tab.go_to(0, 0, end=(2, 1))
    tab.copy()
    small = QGuiApplication.clipboard().text()
    monkeypatch.setattr(sheet_tab, "BACKGROUND_CELLS", -1)
    QGuiApplication.clipboard().clear()
    tab.copy()
    assert QGuiApplication.clipboard().text() == small
    before = {key: c.value for key, c in tab.sheet.cells.items()}
    tab.clear_selection()
    assert all(tab.sheet.cells.get((r, c)) is None or tab.sheet.cells[(r, c)].value is None
               for r in range(3) for c in range(2))
    tab.undo_stack.undo()
    assert {key: c.value for key, c in tab.sheet.cells.items()} == before
    tab.go_to(5, 0)
    tab.paste()
    assert value(tab, "A6") == before.get((0, 0))
    tab.undo_stack.undo()
    assert value(tab, "A6") is None


def test_large_changes_keep_the_window_responsive(tmp_path):
    # 25 000 rows, 50 000 formulas: inserting a row, Undo, clearing every
    # cell, copying them and saving each took the window away for one to
    # several seconds. The bound is generous: a busy machine stays far
    # below it, a change made on the window's thread does not.
    # What earlier tests left on the heap is set aside: a collection of it,
    # due by chance while this one measures, would be counted as this
    # sheet's stall (a whole suite leaves millions of objects).
    gc.collect()
    gc.freeze()
    path = tmp_path / "Big.xlsx"
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("Data")
    rows = 25_000
    for r in range(1, rows + 1):
        sheet.append([r, f"=A{r}*2", f"=B{r}+1", f"Item {r}"])
    sheet.append([None, f"=SUM(B1:B{rows})"])
    book.save(path)
    tab = SheetTab(str(path), background=False)
    gaps, last = [], [time.monotonic()]

    def tick():
        now = time.monotonic()
        gaps.append(now - last[0])
        last[0] = now

    timer = QTimer()
    timer.setInterval(10)
    timer.timeout.connect(tick)
    try:
        tab.resize(900, 600)
        tab.show()
        QTest.qWait(50)
        timer.start()

        def worst(action) -> float:
            QTest.qWait(30)
            gaps.clear()
            last[0] = time.monotonic()
            action()
            QTest.qWait(40)
            return max(gaps)

        stalls = {
            "insert": worst(lambda: tab.insert_rows(0)),
            "undo": worst(tab.undo_stack.undo),
            "select all": worst(tab.grid.selectAll),
            "copy": worst(tab.copy),
            "clear": worst(tab.clear_selection),
            "undo clear": worst(tab.undo_stack.undo),
            "save": worst(lambda: tab.save()),
        }
        assert tab.sheet.cells[(rows, 1)].value == sum(2 * r for r in range(1, rows + 1))
        assert all(stall < 0.5 for stall in stalls.values()), stalls
    finally:
        timer.stop()
        tab.release()
        tab.deleteLater()
        gc.unfreeze()


def test_steps_are_never_taken_from_under_a_running_undo(every_change_is_large, tmp_path):
    # Save As converting the file empties the steps: never while an Undo
    # of a large change is still under way inside them.
    path = tmp_path / "Plan.ods"
    realistic_ods(path)
    tab = SheetTab(str(path))
    try:
        tab.show()
        QApplication.processEvents()
        assert tab.commit_text(0, 0, "5")
        kept = []

        def save_as_xlsx():
            assert tab.is_busy()
            assert tab.save_as(str(tmp_path / "Plan.xlsx")).ok
            kept.append(tab.undo_stack.count())

        QTimer.singleShot(0, save_as_xlsx)
        tab.undo_stack.undo()           # a large change: Save As runs meanwhile
        assert kept and kept[0] > 0     # the steps stayed while the Undo ran
        for _ in range(50):
            if tab.undo_stack.count() == 0:
                break
            QTest.qWait(20)
        assert tab.undo_stack.count() == 0 and not tab.is_busy()
        assert tab._file_path.endswith("Plan.xlsx")
    finally:
        tab.release()
        tab.deleteLater()

# --------------------------------------------------------------------------- #
# Column widths in Excel's unit, numbers that do not fit                      #
# --------------------------------------------------------------------------- #

def _budget(path):
    from openpyxl.styles import Font
    book = openpyxl.Workbook()
    sheet = book.active
    for row, (q3, q4) in enumerate([(135, 135), (24, 0), (0, 480), (300, 150), (90, 120)], 2):
        sheet.cell(row, 2, q3)
        sheet.cell(row, 3, q4)
        sheet.cell(row, 4, f"=B{row}+C{row}")
    for col in "BCD":
        sheet[f"{col}8"] = f"=SUM({col}2:{col}6)"
        sheet[f"{col}8"].font = Font(bold=True)
    for row in range(2, 9):
        for col in range(2, 6):
            sheet.cell(row, col).number_format = '#,##0.00 "€"'
    sheet.column_dimensions["E"].width = 4
    sheet["E8"] = "=D8"
    sheet["F8"] = "A long line of words"
    sheet.column_dimensions["F"].width = 4
    sheet.column_dimensions["G"].width = 10.5
    book.save(path)


def _shown(tab, ref):
    """What the grid draws in ``ref``, and the room it has there."""
    row, col = sf.parse_cell_name(ref)
    index = tab.model.index(row, col)
    option = QStyleOptionViewItem()
    tab.grid.initViewItemOption(option)
    option.rect = tab.grid.visualRect(index)
    tab.grid.itemDelegate().initStyleOption(option, index)
    style = tab.grid.style()
    margin = style.pixelMetric(QStyle.PixelMetric.PM_FocusFrameHMargin, None, tab.grid) + 1
    room = style.subElementRect(QStyle.SubElement.SE_ItemViewItemText, option,
                                tab.grid).width() - 2 * margin
    return option, room


@pytest.fixture
def budget(tmp_path):
    path = tmp_path / "Budget 2026.xlsx"
    _budget(path)
    tab = SheetTab(str(path))
    tab._style, tab._locale = ENGLISH, US
    tab.resize(900, 600)
    tab.show()
    QApplication.processEvents()
    yield tab
    tab.release()
    tab.deleteLater()


def test_a_bold_total_at_excels_default_width_shows_in_full(budget):
    # Excel holds "1,434.00 €" in a column of the default width; so does
    # the grid, whose digits are wider than Calibri's.
    option, room = _shown(budget, "D8")
    assert option.text == "1,434.00 €" and option.font.bold()
    assert QFontMetrics(option.font).horizontalAdvance(option.text) <= room


def test_a_number_too_wide_for_its_column_shows_hashes(budget):
    option, room = _shown(budget, "E8")
    assert option.text and set(option.text) == {"#"}
    assert QFontMetrics(option.font).horizontalAdvance(option.text) <= room
    words, _room = _shown(budget, "F8")
    assert words.text == "A long line of words"         # text is shortened as drawn


def test_column_widths_keep_the_files_unit_exactly(budget):
    path = budget._file_path
    assert budget.sheet.widths[6] == 10.5 * 7 + 5       # read exactly, no rounding
    budget.grid.horizontalHeader().resizeSection(1, 150)
    assert budget.grid.columnWidth(1) == 150
    assert budget.save().ok
    saved = openpyxl.load_workbook(path).active.column_dimensions
    width = saved["B"].width
    assert saved["G"].width == 10.5
    again = SheetTab(path)
    try:
        again.show()
        QApplication.processEvents()
        assert again.grid.columnWidth(1) == 150          # as dragged
        again.grid.horizontalHeader().resizeSection(1, 150)
        again.commit_text(0, 0, "x")
        assert again.save().ok
        assert openpyxl.load_workbook(path).active.column_dimensions["B"].width == width
    finally:
        again.release()
        again.deleteLater()
